"""Local experimental HWP-to-DOCX conversion.

본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.
"""

import argparse
import base64
from dataclasses import asdict
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from urllib.parse import urlsplit

from hwp3_parser import SIGNATURE as HWP3_SIGNATURE, validate_header as validate_hwp3_header
from hwp_limits import HwpLimits, MIB, add_limit_arguments, positive_integer

import olefile
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.section import _Header, _Footer
from docx.shared import Inches, Pt, RGBColor


NOTICE = "본 제품은 한컴의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다."
WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PAGE_ATTRIBUTES = (
    ("page_width", "width"), ("page_height", "height"),
    ("top_margin", "top-offset"), ("bottom_margin", "bottom-offset"),
    ("left_margin", "left-offset"), ("right_margin", "right-offset"),
    ("header_distance", "header-offset"), ("footer_distance", "footer-offset"),
)


def page_orientation(page):
    landscape = (int(page.get("flags")) & 1 if page.get("flags") is not None
                 else int(page.get("width", "0")) > int(page.get("height", "0")))
    return WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT


def word_property(parent, name, **attributes):
    element = OxmlElement("w:" + name)
    for key, value in attributes.items():
        element.set(qn("w:" + key), str(value))
    parent.append(element)
    return element


def story_elements(node):
    """Elements in one document story, excluding embedded header/footer stories."""
    if node.tag in ("Header", "Footer"):
        return
    yield node
    for child in node:
        yield from story_elements(child)


def header_has_content(item):
    return (any(source_text(paragraph) for paragraph in item.findall("Paragraph"))
            or any(node.tag in ("TableControl", "GShapeObjectControl") for node in item.iter()))


def header_footer_bindings(source_section, different_pages, allow_empty_late=False):
    definitions = {}
    for index, paragraph in enumerate(source_section.findall("ColumnSet/Paragraph")):
        for item in paragraph.iter():
            if item.tag not in ("Header", "Footer"):
                continue
            if index:
                if allow_empty_late and not header_has_content(item):
                    continue
                raise ValueError("Headers/footers defined after the first section paragraph are unsupported.")
            selection = item.get("apply", "both")
            if selection not in ("both", "odd", "even"):
                raise ValueError("Invalid header/footer page selection.")
            key = (item.tag, selection)
            if key in definitions:
                raise ValueError("Repeated header/footer definitions for the same pages are unsupported.")
            definitions[key] = item
    bindings = {}
    for kind in ("Header", "Footer"):
        name = kind.lower()
        both = definitions.get((kind, "both"))
        odd = definitions.get((kind, "odd"))
        even = definitions.get((kind, "even"))
        if odd is not None or both is not None:
            bindings[name] = odd if odd is not None else both
        if different_pages and (even is not None or both is not None):
            bindings["even_page_" + name] = even if even is not None else both
    return bindings


def hyperlink_destination(item):
    target = item.get("target", "")
    if not target or any(ord(char) < 32 for char in target):
        return None
    try:
        parsed = urlsplit(target)
        if parsed.scheme.lower() in ("http", "https", "ftp") and parsed.hostname:
            return target
        if parsed.scheme.lower() == "mailto" and parsed.path:
            return target
    except ValueError:
        pass
    return None


def read_hwp(source, recover_text=False, limits=None, reader_timeout=60, asset_directory=None):
    limits = limits if limits is not None else HwpLimits()
    if type(reader_timeout) is not int or reader_timeout <= 0:
        raise ValueError("Reader timeout must be a positive integer number of seconds.")
    if source.suffix.lower() != ".hwp":
        raise ValueError("This prototype supports binary .hwp files only, not HWPX.")
    limits.check_input(source)
    with source.open("rb") as stream:
        prefix = stream.read(158)
    if prefix[:30] == HWP3_SIGNATURE:
        validate_hwp3_header(prefix)
    else:
        with olefile.OleFileIO(str(source)) as container:
            header_size = container.get_size("FileHeader")
            if header_size > limits.max_stream_mb * MIB:
                raise ValueError("HWP header exceeds the size limit (--max-stream-mb).")
            if header_size > limits.max_expanded_mb * MIB:
                raise ValueError("HWP header exceeds the size limit (--max-expanded-mb).")
            header = container.openstream("FileHeader").read(40)
        if not header.startswith(b"HWP Document File") or len(header) < 40:
            raise ValueError("Not a supported HWP document.")
        if header[35] != 5:
            raise ValueError("Only binary HWP 3 and HWP 5 files are supported.")
        flags = struct.unpack_from("<I", header, 36)[0]
        if flags & ((1 << 1) | (1 << 2) | (1 << 4) | (1 << 8) | (1 << 10)):
            raise ValueError("Password, distribution, or DRM protected files are not converted.")
    command = [sys.executable, str(Path(__file__).with_name("hwp_reader.py")), str(source),
               *limits.arguments()]
    if recover_text:
        command.append("--recover-text")
    if asset_directory is not None:
        command.extend(("--asset-directory", str(asset_directory)))
    # Spool the worker's XML to disk instead of retaining another complete copy
    # in captured stdout alongside the parsed document and embedded images.
    with tempfile.TemporaryFile() as intermediate:
        try:
            result = subprocess.run(command, stdout=intermediate, stderr=subprocess.PIPE,
                                    timeout=reader_timeout)
        except subprocess.TimeoutExpired as error:
            raise ValueError(f"HWP reader exceeded {reader_timeout} seconds; "
                             "use --reader-timeout to raise the timeout.") from error
        if result.returncode:
            detail = result.stderr.decode("utf-8", errors="replace").strip()
            raise ValueError("The native HWP reader failed; no output was created. " + detail)
        intermediate.seek(0)
        root = ET.parse(intermediate).getroot()
    return root, json.loads(root.get("parser-warnings", "[]"))


class Converter:
    def __init__(self, root, warnings, font_map=None, preserve_source_pages=False, asset_directory=None):
        self.root = root
        self.preserve_source_pages = preserve_source_pages
        self.font_map = font_map or {}
        self.asset_directory = Path(asset_directory).resolve() if asset_directory is not None else None
        self.document = Document()
        self.warnings = set(warnings)
        self.tables = 0
        self.images = 0
        self.assets = {node.get("storage-id"): node for node in root.iter("BinDataEmbedding")}
        self.char_shapes = list(root.iter("CharShape"))
        self.para_shapes = list(root.iter("ParaShape"))
        self.tab_defs = list(root.iter("TabDef"))
        self.borders = list(root.iter("BorderFill"))
        mappings = root.find("DocInfo/IdMappings")
        faces = list(root.iter("FaceName"))
        self.fonts = {}
        offset = 0
        for language in ("ko", "en", "cn", "jp", "other", "symbol", "user"):
            count = int(mappings.get(language + "-fonts", "0")) if mappings is not None else 0
            self.fonts[language] = [face.get("name") for face in faces[offset:offset + count]]
            offset += count

    def style_run(self, run, shape_id, language="ko"):
        if not self.char_shapes:
            return
        shape = self.char_shapes[int(shape_id)]
        run.bold = shape.get("bold") == "1"
        run.italic = shape.get("italic") == "1"
        run.underline = shape.get("underline", "none") != "none"
        if shape.get("superscript") == "1":
            run.font.superscript = True
        if shape.get("subscript") == "1":
            run.font.subscript = True
        run.font.size = Pt(int(shape.get("basesize", "1000")) / 100)
        run.font.color.rgb = RGBColor.from_string(shape.get("text-color", "#000000").lstrip("#"))
        spacing = shape.find("LetterSpacing")
        if spacing is not None:
            percentage = int(spacing.get(language, spacing.get("ko", "0")))
            word_property(run._element.get_or_add_rPr(), "spacing", val=round(run.font.size.pt * percentage / 100 * 20))
        expansion = shape.find("LetterWidthExpansion")
        if expansion is not None:
            word_property(run._element.get_or_add_rPr(), "w", val=expansion.get(language, expansion.get("ko", "100")))
        face = shape.find("FontFace")
        if face is not None and self.fonts["ko"] and self.fonts["en"]:
            korean = self.fonts["ko"][int(face.get("ko", "0"))]
            latin = self.fonts["en"][int(face.get("en", "0"))]
            for name in (korean, latin):
                if name in self.font_map:
                    self.warnings.add("Font substitution: " + name + " -> " + self.font_map[name])
            korean = self.font_map.get(korean, korean)
            latin = self.font_map.get(latin, latin)
            run.font.name = latin
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), korean)

    def style_paragraph(self, source, paragraph):
        settings = paragraph.paragraph_format
        settings.space_before = Pt(0)
        settings.space_after = Pt(0)
        settings.page_break_before = source.get("new-page") == "1"
        if not self.para_shapes:
            return
        shape = self.para_shapes[int(source.get("parashape-id", "0"))]
        paragraph.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "center": WD_ALIGN_PARAGRAPH.CENTER, "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "distribute": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
        }.get(shape.get("align"), WD_ALIGN_PARAGRAPH.LEFT)
        indent = int(shape.get("indent", "0")) / 100
        settings.left_indent = Pt(int(shape.get("doubled-margin-left", "0")) / 200 + max(0, -indent))
        settings.right_indent = Pt(int(shape.get("doubled-margin-right", "0")) / 200)
        settings.first_line_indent = Pt(indent)
        settings.space_before = Pt(int(shape.get("doubled-margin-top", "0")) / 200)
        settings.space_after = Pt(int(shape.get("doubled-margin-bottom", "0")) / 200)
        if shape.get("linespacing-type") == "ratio":
            lines = source.findall("LineSeg")
            if lines and all(line.get("height") is not None for line in lines):
                settings.line_spacing = Pt(max(int(line.get("height", "1000")) + int(line.get("space-below", "0")) for line in lines) / 100)
            else:
                settings.line_spacing = int(shape.get("linespacing", "100")) / 100
        elif shape.get("linespacing-type") == "fixed":
            settings.line_spacing = Pt(int(shape.get("linespacing")) / 100)
        else:
            self.warnings.add("Non-proportional paragraph line spacing is approximated.")
        settings.keep_with_next = shape.get("with-next-paragraph") == "1"
        settings.keep_together = shape.get("protect") == "1"
        settings.widow_control = shape.get("protect-single-line") == "1"
        settings.page_break_before = source.get("new-page") == "1" or shape.get("start-new-page") == "1"
        tab_id = int(shape.get("tabdef-id", "0"))
        if tab_id < len(self.tab_defs):
            alignments = {"left": WD_TAB_ALIGNMENT.LEFT, "right": WD_TAB_ALIGNMENT.RIGHT,
                          "center": WD_TAB_ALIGNMENT.CENTER, "decimal": WD_TAB_ALIGNMENT.DECIMAL}
            for tab in self.tab_defs[tab_id].findall("Array/Tab"):
                settings.tab_stops.add_tab_stop(Pt(int(tab.get("pos", "0")) / 100),
                    alignments.get(tab.get("kind"), WD_TAB_ALIGNMENT.LEFT),
                    WD_TAB_LEADER.DOTS if tab.get("fill-type") == "3" else WD_TAB_LEADER.SPACES)
        if shape.get("head-shape", "none") != "none":
            self.warnings.add("Automatic numbering/bullets are not yet mapped.")

    def style_cell(self, source, cell, body=None):
        cell.vertical_alignment = {
            "top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "middle": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
            "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
        }.get(source.get("valign"), WD_CELL_VERTICAL_ALIGNMENT.TOP)
        properties = cell._tc.get_or_add_tcPr()
        margins = word_property(properties, "tcMar")
        padding = source
        lines = source.findall("Paragraph/LineSeg")
        if body is not None and lines:
            expected_width = int(source.get("width", "0")) - int(body.get("padding-left", "0")) - int(body.get("padding-right", "0"))
            if all(abs(int(line.get("width", "0")) - expected_width) <= 2 for line in lines):
                padding = body
        for side in ("top", "left", "bottom", "right"):
            word_property(margins, side, w=round(int(padding.get("padding-" + side, "0")) / 5), type="dxa")
        border_id = int(source.get("borderfill-id", "0"))
        if not border_id or border_id > len(self.borders):
            return
        fill = self.borders[border_id - 1]
        borders = word_property(properties, "tcBorders")
        strokes = {"none": "nil", "solid": "single", "double-2": "double", "dashed": "dashed", "dotted": "dotted"}
        for border in fill.findall("Border"):
            side = border.get("attribute-name")
            if side not in ("top", "left", "bottom", "right"):
                continue
            width = float(border.get("width", "0.1mm").removesuffix("mm"))
            stroke = border.get("stroke-type", "none")
            if stroke not in strokes:
                self.warnings.add("Some border styles are approximated.")
            word_property(borders, side, val=strokes.get(stroke, "single"),
                          sz=max(2, round(width * 72 / 25.4 * 8)), color=border.get("color", "#000000").lstrip("#"))
        color = fill.find("FillColorPattern")
        if color is not None:
            word_property(properties, "shd", val="clear", fill=color.get("background-color", "#ffffff").lstrip("#"))

    def append_text(self, item, paragraph):
        run = paragraph.add_run(item.text or "")
        self.style_run(run, item.get("charshape-id", "0"), item.get("lang", "ko"))
        for child in item:
            if child.tag == "Tab":
                run.add_tab()
            if child.tail:
                run.add_text(child.tail)

    def append_control(self, item, paragraph):
        name = item.get("name")
        if name == "LINE_BREAK":
            paragraph.add_run().add_break()
        elif name == "TAB":
            paragraph.add_run().add_tab()
        elif name == "FIXWIDTH_SPACE":
            paragraph.add_run("\u2007")
            self.warnings.add("Fixed-width spaces are approximated with Unicode figure spaces.")
        elif name != "PARAGRAPH_BREAK":
            self.warnings.add("Unsupported text control: " + str(name))

    def hyperlink(self, item, paragraph):
        target = hyperlink_destination(item)
        before = set(paragraph._p)
        for child in item:
            if child.tag == "Text":
                self.append_text(child, paragraph)
            elif child.tag == "ControlChar":
                self.append_control(child, paragraph)
            elif child.tag == "GShapeObjectControl":
                self.drawing(child, paragraph)
            else:
                raise ValueError("Unsupported content inside a hyperlink.")
        if target is None:
            self.warnings.add("A hyperlink has no supported web, email, or FTP target; its label remains as text.")
            return
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
        for child in list(paragraph._p):
            if child not in before:
                link.append(child)
        paragraph._p.append(link)

    def drawing(self, item, paragraph):
        seen_paragraph = False
        def nodes(parent):
            yield parent
            if parent.tag != "FieldHyperLink":
                for child in parent:
                    yield from nodes(child)
        for node in nodes(item):
            if node.tag == "Paragraph":
                if seen_paragraph:
                    paragraph.add_run().add_break()
                seen_paragraph = True
            elif node.tag == "Text":
                self.append_text(node, paragraph)
            elif node.tag == "ControlChar":
                self.append_control(node, paragraph)
            elif node.tag == "FieldHyperLink":
                self.hyperlink(node, paragraph)
            elif node.tag == "ShapePicture":
                picture = node.find("PictureInfo")
                identifier = int(picture.get("bindata-id"))
                asset = self.assets.get(f"BIN{identifier:04X}")
                if asset is None:
                    raise ValueError(f"Embedded image {identifier} is unavailable; no output published.")
                if asset.get("file") is not None and self.asset_directory is not None:
                    image_path = (self.asset_directory / asset.get("file")).resolve()
                    if image_path.parent != self.asset_directory:
                        raise ValueError("Invalid extracted image path.")
                    image_input = str(image_path)
                elif asset.get("inline") == "true":
                    image_input = BytesIO(base64.b64decode("".join((asset.text or "").split()), validate=True))
                else:
                    raise ValueError(f"Embedded image {identifier} is unavailable; no output published.")
                width = int(item.get("width", "0"))
                height = int(item.get("height", "0"))
                if width <= 0 or height <= 0:
                    raise ValueError("Image dimensions must be positive.")
                try:
                    paragraph.add_run().add_picture(image_input, width=Inches(width / 7200), height=Inches(height / 7200))
                except Exception as error:
                    raise ValueError(f"Unable to insert embedded image {identifier} ({asset.get('ext')}).") from error
                self.images += 1
                paragraph.paragraph_format.line_spacing = 1.0
        if list(item.iter("TextboxParagraphList")):
            self.warnings.add("Text boxes are flattened to editable inline text; box layout is not preserved.")
        self.warnings.add("Drawings are flattened: floating positions, crops, rotations, and effects are not preserved.")

    def paragraph(self, source, target, page_break_before=False):
        paragraph = target.add_paragraph()
        self.style_paragraph(source, paragraph)
        if page_break_before:
            paragraph.paragraph_format.page_break_before = True
        follows_table = False
        for line in source.findall("LineSeg"):
            for item in line:
                if item.tag == "Text":
                    follows_table = False
                    self.append_text(item, paragraph)
                elif item.tag == "GShapeObjectControl":
                    follows_table = False
                    self.drawing(item, paragraph)
                elif item.tag == "FieldHyperLink":
                    follows_table = False
                    self.hyperlink(item, paragraph)
                elif item.tag == "TableControl":
                    # Runs can contain pictures or breaks without any text.
                    if not paragraph._element.xpath(".//w:r/*[not(self::w:rPr)]"):
                        if paragraph.paragraph_format.page_break_before:
                            self.compact_paragraph(paragraph)
                            paragraph.paragraph_format.keep_with_next = True
                        else:
                            paragraph._element.getparent().remove(paragraph._element)
                    self.table(item, target, paragraph.alignment)
                    paragraph = target.add_paragraph()
                    self.style_paragraph(source, paragraph)
                    paragraph.paragraph_format.page_break_before = False
                    follows_table = True
                elif item.tag == "ControlChar":
                    self.append_control(item, paragraph)
                elif item.tag == "ColumnsDef":
                    if item.get("count") != "1":
                        self.warnings.add("Multiple columns are flattened.")
                elif item.tag in ("Header", "Footer"):
                    continue
                else:
                    self.warnings.add("Unsupported paragraph element: " + item.tag)
        if follows_table and not paragraph.runs:
            self.compact_paragraph(paragraph)
        elif not paragraph.runs:
            last = paragraph._element.getparent()
            if last is not None:
                line = source.find("LineSeg")
                if line is not None:
                    paragraph.add_run().font.size = Pt(int(line.get("height-text", "1000")) / 100)

    @staticmethod
    def compact_paragraph(paragraph):
        settings = paragraph.paragraph_format
        settings.line_spacing = Pt(1)
        settings.space_before = Pt(0)
        settings.space_after = Pt(0)
        paragraph.add_run().font.size = Pt(1)

    def table(self, source, target, paragraph_alignment=None):
        body = source.find("TableBody")
        rows, cols = int(body.get("rows")), int(body.get("cols"))
        if rows < 1 or cols < 1 or rows * cols > 10000:
            raise ValueError("Table dimensions exceed prototype limits.")
        width = Inches(int(source.get("width", "43200")) / 7200)
        table = (target.add_table(rows, cols, width) if isinstance(target, (_Header, _Footer))
                 else target.add_table(rows, cols))
        table.autofit = False
        if source.get("inline", "1") == "1":
            table.alignment = {
                WD_ALIGN_PARAGRAPH.CENTER: WD_TABLE_ALIGNMENT.CENTER,
                WD_ALIGN_PARAGRAPH.RIGHT: WD_TABLE_ALIGNMENT.RIGHT,
            }.get(paragraph_alignment, WD_TABLE_ALIGNMENT.LEFT)
        preferred_width = table._tbl.tblPr.find(qn("w:tblW"))
        preferred_width.set(qn("w:type"), "dxa")
        preferred_width.set(qn("w:w"), str(round(width.twips)))
        if source.get("inline", "1") != "1":
            self.warnings.add("Floating tables are converted to inline tables.")
        self.tables += 1
        cells = body.findall("TableRow/TableCell")
        for item in cells:
            row, col = int(item.get("row")), int(item.get("col"))
            rowspan, colspan = int(item.get("rowspan", "1")), int(item.get("colspan", "1"))
            cell = table.cell(row, col)
            if rowspan > 1 or colspan > 1:
                cell = cell.merge(table.cell(row + rowspan - 1, col + colspan - 1))
            cell.width = Inches(int(item.get("width", "7200")) / 7200)
            if colspan == 1:
                table.columns[col].width = cell.width
            self.style_cell(item, cell, body)
            if rowspan == 1:
                table.rows[row].height = Inches(int(item.get("height", "0")) / 7200)
                table.rows[row].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for paragraph in list(cell.paragraphs):
                cell._tc.remove(paragraph._element)
            for paragraph in item.findall("Paragraph"):
                self.paragraph(paragraph, cell)
                lines = paragraph.findall("LineSeg")
                if lines and len(lines) == 1 and lines[0].get("height") is not None and not paragraph.findall(".//TableControl"):
                    cell.paragraphs[-1].paragraph_format.line_spacing = Pt(int(lines[0].get("height", "1000")) / 100)
            if not cell.paragraphs:
                cell.add_paragraph()
        if source.find("TableCaption") is not None:
            self.warnings.add("Table captions are not yet supported.")

    def convert(self):
        if self.preserve_source_pages and self.root.get("parser") == "native-hwp3":
            self.warnings.add("--preserve-source-pages inference is only available for HWP 5; HWP 3 retains explicit paragraph page breaks.")
            self.preserve_source_pages = False
        sections = self.root.findall("BodyText/SectionDef")
        if not sections:
            raise ValueError("The document has no sections.")
        if self.preserve_source_pages and any(node.get("count", "1") != "1"
                                              for source_section in sections
                                              for node in source_section.iter("ColumnsDef")):
            raise ValueError("Source page inference is only supported for single-column documents.")
        self.warnings.add("Font availability, pagination, and exact layout require review in Word.")
        if len(sections) > 1:
            self.warnings.add("Section boundaries start new Word pages; odd/even starts and section numbering are not preserved.")
        different_pages = any(item.get("apply") in ("odd", "even")
                              for item in self.root.iter() if item.tag in ("Header", "Footer"))
        self.document.settings.odd_and_even_pages_header_footer = different_pages
        allow_empty_late = not any(header_has_content(item) for item in self.root.iter()
                                   if item.tag in ("Header", "Footer"))
        for index, source_section in enumerate(sections):
            section = (self.document.sections[0] if index == 0
                       else self.document.add_section(WD_SECTION_START.NEW_PAGE))
            page = source_section.find("PageDef")
            if page is not None:
                for target, attribute in PAGE_ATTRIBUTES:
                    setattr(section, target, Inches(int(page.get(attribute, "0")) / 7200))
                section.orientation = page_orientation(page)
            for name, source_story in header_footer_bindings(source_section, different_pages, allow_empty_late).items():
                target_story = getattr(section, name)
                target_story.is_linked_to_previous = False
                for child in list(target_story._element):
                    target_story._element.remove(child)
                for source_paragraph in source_story.findall("Paragraph"):
                    self.paragraph(source_paragraph, target_story)
                if not target_story.paragraphs or target_story._element[-1].tag != qn("w:p"):
                    target_story.add_paragraph()
            previous_y = None
            for paragraph in source_section.findall("ColumnSet/Paragraph"):
                lines = paragraph.findall("LineSeg")
                first_y = int(lines[0].get("y", "0")) if lines else None
                source_page_break = (self.preserve_source_pages and first_y is not None
                                     and previous_y is not None and first_y < previous_y)
                self.paragraph(paragraph, self.document, page_break_before=source_page_break)
                if lines:
                    previous_y = int(lines[-1].get("y", "0"))
        if self.preserve_source_pages:
            self.warnings.add("Page breaks inferred from stored line positions; not a visual fidelity guarantee.")
        return self.document


def source_text(node):
    """Comparable text, including inline controls and flattened textbox boundaries."""
    if node.tag in ("Header", "Footer"):
        return ""
    if node.tag == "Text":
        return (node.text or "") + "".join(
            ("\t" if child.tag == "Tab" else "") + (child.tail or "") for child in node)
    if node.tag == "ControlChar":
        return {"TAB": "\t", "LINE_BREAK": "\n", "FIXWIDTH_SPACE": "\u2007"}.get(node.get("name"), "")
    if node.tag == "GShapeObjectControl":
        chunks = []
        seen_paragraph = False
        for child in node.iter():
            if child.tag == "Paragraph":
                if seen_paragraph:
                    chunks.append("\n")
                seen_paragraph = True
            elif child.tag in ("Text", "ControlChar"):
                chunks.append(source_text(child))
        return "".join(chunks)
    return "".join(source_text(child) for child in node)


def word_text(document):
    content = []
    # A w:tab in paragraph properties defines a tab stop, not document text.
    for node in document.findall(".//" + WORD_NS + "r/*"):
        if node.tag == WORD_NS + "t":
            content.append(node.text or "")
        elif node.tag == WORD_NS + "tab":
            content.append("\t")
        elif node.tag == WORD_NS + "cr" or (node.tag == WORD_NS + "br"
                and node.get(WORD_NS + "type", "textWrapping") == "textWrapping"):
            content.append("\n")
    return "".join(content)


def validate_story(source, document, part):
    if source_text(source) != word_text(document):
        raise ValueError("Text preservation check failed; output is not verified.")
    elements = list(story_elements(source))
    source_tables = [node for node in elements if node.tag == "TableBody"]
    output_tables = list(document.iter(WORD_NS + "tbl"))
    if len(source_tables) != len(output_tables):
        raise ValueError("Table preservation check failed.")
    for source_table, output_table in zip(source_tables, output_tables):
        if len(output_table.findall(WORD_NS + "tr")) != int(source_table.get("rows")):
            raise ValueError("Table row count changed.")
        if len(output_table.find(WORD_NS + "tblGrid")) != int(source_table.get("cols")):
            raise ValueError("Table column count changed.")
    pictures = [node for node in elements if node.tag == "ShapePicture"]
    drawings = list(document.iter("{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"))
    if len(pictures) != len(drawings):
        raise ValueError("Image occurrence count changed.")
    links = [node for node in elements if node.tag == "FieldHyperLink" and hyperlink_destination(node)]
    actual_links = list(document.iter(WORD_NS + "hyperlink"))
    if len(links) != len(actual_links):
        raise ValueError("Hyperlink count changed.")
    for expected, actual in zip(links, actual_links):
        relationship = part.rels.get(actual.get(qn("r:id")))
        if (relationship is None or relationship.reltype != RT.HYPERLINK
                or not relationship.is_external or relationship.target_ref != hyperlink_destination(expected)
                or source_text(expected) != word_text(actual)):
            raise ValueError("Hyperlink destination or label changed.")
    return {"tables": len(output_tables), "images": len(drawings), "hyperlinks": len(links)}


def validate_output(root, output):
    with ZipFile(output) as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
        if archive.testzip() is not None:
            raise ValueError("Generated DOCX archive is damaged.")
    reopened = Document(output)
    counts = validate_story(root.find("BodyText"), document, reopened.part)
    # Preserve the original body-only text hash for existing report consumers.
    actual = "".join(node.text or "" for node in document.iter(WORD_NS + "t"))
    fixed_spaces = sum(node.tag == "ControlChar" and node.get("name") == "FIXWIDTH_SPACE"
                       for node in story_elements(root.find("BodyText")))
    if fixed_spaces:
        actual = actual.replace("\u2007", "")
    source_sections = root.findall("BodyText/SectionDef")
    if len(reopened.sections) != len(source_sections):
        raise ValueError("Section count changed.")
    different_pages = any(item.get("apply") in ("odd", "even")
                          for item in root.iter() if item.tag in ("Header", "Footer"))
    if reopened.settings.odd_and_even_pages_header_footer != different_pages:
        raise ValueError("Odd/even header/footer settings changed.")
    stories = 0
    story_characters = 0
    allow_empty_late = not any(header_has_content(item) for item in root.iter()
                               if item.tag in ("Header", "Footer"))
    for index, (source_section, section) in enumerate(zip(source_sections, reopened.sections)):
        if index and section.start_type != WD_SECTION_START.NEW_PAGE:
            raise ValueError("Section page break changed.")
        page = source_section.find("PageDef")
        if page is not None:
            for target, attribute in PAGE_ATTRIBUTES:
                expected_size = Inches(int(page.get(attribute, "0")) / 7200)
                actual_size = getattr(section, target)
                if actual_size is None or actual_size.twips != expected_size.twips:
                    raise ValueError(f"Section {index + 1} {target} changed.")
            if section.orientation != page_orientation(page):
                raise ValueError(f"Section {index + 1} orientation changed.")
        bindings = header_footer_bindings(source_section, different_pages, allow_empty_late)
        for name in ("header", "footer", "even_page_header", "even_page_footer"):
            if name not in bindings and not getattr(section, name).is_linked_to_previous:
                raise ValueError("Header/footer inheritance changed.")
        for name, source_story in bindings.items():
            target_story = getattr(section, name)
            if target_story.is_linked_to_previous:
                raise ValueError("Explicit header/footer definition is missing.")
            story = ET.Element("Story")
            story.extend(source_story)
            story_counts = validate_story(story, target_story._element, target_story.part)
            for key, value in story_counts.items():
                counts[key] += value
            stories += 1
            story_characters += len(source_text(story))
    return {"text_characters": len(actual), "text_sha256": hashlib.sha256(actual.encode()).hexdigest(),
            **counts, "text_preserved": True, "table_dimensions_preserved": True,
            "image_occurrences_preserved": True, "hyperlink_destinations_preserved": True,
            "header_footer_stories": stories, "header_footer_text_characters": story_characters,
            "header_footer_content_preserved": True,
            "sections": len(source_sections), "section_page_settings_preserved": True,
            "fixed_width_spaces_approximated": fixed_spaces}


def publish_output(candidate, output, report):
    """Publish complete files without replacing either destination."""
    report_path = output.with_suffix(".report.json")
    staged_report = candidate.with_suffix(".report.json")
    staged_report.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    try:
        # Both staged files are on the destination filesystem. Linking the report
        # first also reserves this output pair against concurrent conversions.
        os.link(staged_report, report_path)
        try:
            os.link(candidate, output)
        except BaseException:
            # Remove only our report if publication of the document failed.
            if report_path.exists() and report_path.samefile(staged_report):
                report_path.unlink()
            raise
    except FileExistsError:
        raise ValueError("Output or report already exists; choose a new name.") from None


def main():
    parser = argparse.ArgumentParser(description=NOTICE)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--font-map", type=Path, help="Optional JSON mapping from original font families to installed alternatives.")
    parser.add_argument("--recover-text", action="store_true", help="Replace malformed HWP 5 UTF-16 with U+FFFD and report each replacement; original file is unchanged.")
    parser.add_argument("--preserve-source-pages", action="store_true", help="Infer top-level page breaks from stored HWP 5 line positions; intended for single-column documents.")
    add_limit_arguments(parser)
    parser.add_argument("--reader-timeout", type=positive_integer, default=60,
                        help="Reader timeout in seconds; default: 60.")
    args = parser.parse_args()
    output = args.output
    if output.suffix.lower() != ".docx":
        parser.error("Output must have a .docx extension.")
    report_path = output.with_suffix(".report.json")
    if output.exists() or report_path.exists():
        parser.error("Output or report already exists; choose a new name.")
    try:
        limits = HwpLimits.from_arguments(args)
        with tempfile.TemporaryDirectory(prefix="hwp-assets-") as asset_directory:
            root, warnings = read_hwp(args.source, recover_text=args.recover_text, limits=limits,
                                     reader_timeout=args.reader_timeout, asset_directory=asset_directory)
            font_map = json.loads(args.font_map.read_text(encoding="utf-8")) if args.font_map else {}
            if not isinstance(font_map, dict) or not all(isinstance(key, str) and isinstance(value, str) and value for key, value in font_map.items()):
                raise ValueError("Font map must be a JSON object mapping font names to nonempty strings.")
            converter = Converter(root, warnings, font_map, preserve_source_pages=args.preserve_source_pages,
                                  asset_directory=asset_directory)
            document = converter.convert()
            output.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(dir=output.parent) as temporary:
                candidate = Path(temporary) / "candidate.docx"
                document.save(candidate)
                report = validate_output(root, candidate)
                report["warnings"] = sorted(converter.warnings)
                report["parser"] = root.get("parser", "unknown")
                report["text_replacements"] = json.loads(root.get("text-replacements", "[]"))
                report["text_comparison_basis"] = "Recovered parser text" if report["text_replacements"] else "Parser text"
                report["visual_layout_verified"] = False
                report["reader_timeout_seconds"] = args.reader_timeout
                if root.get("parser") == "native-hwp5":
                    report["hwp5_limits"] = {**asdict(limits), "size_unit": "MiB"}
                publish_output(candidate, output, report)
    except (ValueError, OSError, subprocess.TimeoutExpired, ET.ParseError) as error:
        parser.exit(1, str(error) + "\n")
    print(json.dumps(report, indent=2))
    print("Created:", output)


if __name__ == "__main__":
    main()
