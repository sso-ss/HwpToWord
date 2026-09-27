"""HWP 5 resource budgets and multi-section conversion regressions."""

from io import BytesIO
import json
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch
import xml.etree.ElementTree as ET
import zlib

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.shared import Inches

from convert_hwp import Converter, read_hwp, validate_output
from hwp_limits import HwpLimits, MIB
from hwp_parser import HwpParser, record_tree
from test_convert_hwp import paragraph, sample_root
from test_hwp_parser import control_text, paragraph_records, record


def compound_file(streams, assets=()):
    """Small standard CFB v3 fixture using regular streams (at least 4096 bytes)."""
    end, free, fat_sector = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD
    # Root: FileHeader -> DocInfo -> BodyText. BodyText owns SectionN streams.
    entries = [("Root Entry", 5, free, 1, end, 0),
               ("FileHeader", 2, 2, free, 0, 0),
               ("DocInfo", 2, 3, free, 0, 0),
               ("BodyText", 1, free, 4 if len(streams) > 2 else free, end, 0)]
    for index, (name, data) in enumerate(streams[2:], 4):
        entries.append((name, 2, index + 1 if index < len(streams) + 1 else free,
                        free, 0, 0))
    stream_entries = [1, 2, *range(4, len(entries))]
    if assets:
        storage_index = len(entries)
        entries[3] = ("BodyText", 1, storage_index, entries[3][3], end, 0)
        entries.append(("BinData", 1, free, storage_index + 1, end, 0))
        for index, (name, data) in enumerate(assets):
            entry_index = len(entries)
            entries.append((name, 2, entry_index + 1 if index + 1 < len(assets) else free, free, 0, 0))
            stream_entries.append(entry_index)
    directory_sectors = (len(entries) + 3) // 4
    sectors = [b""] * (directory_sectors + 1)
    fat = [index + 1 if index + 1 < directory_sectors else end
           for index in range(directory_sectors)] + [fat_sector]
    for (name, data), entry_index in zip([*streams, *assets], stream_entries):
        assert len(data) >= 4096
        first = len(sectors)
        chunks = [data[offset:offset + 512].ljust(512, b"\0") for offset in range(0, len(data), 512)]
        sectors.extend(chunks)
        fat.extend([first + index + 1 if index + 1 < len(chunks) else end for index in range(len(chunks))])
        old = entries[entry_index]
        entries[entry_index] = (*old[:4], first, len(data))
    assert len(fat) <= 128
    directory = bytearray()
    for name, kind, right, child, start, size in entries:
        data = bytearray(128)
        encoded = (name + "\0").encode("utf-16le")
        data[:len(encoded)] = encoded
        struct.pack_into("<HBBIII", data, 64, len(encoded), kind, 1, free, right, child)
        struct.pack_into("<IQ", data, 116, start, size)
        directory.extend(data)
    directory = directory.ljust(directory_sectors * 512, b"\0")
    sectors[:directory_sectors] = [directory[i:i + 512] for i in range(0, len(directory), 512)]
    sectors[directory_sectors] = struct.pack("<128I", *(fat + [free] * (128 - len(fat))))
    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<5H", header, 24, 0x003E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<9I", header, 40, 0, 1, 0, 0, 4096, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, directory_sectors, *([free] * 108))
    return bytes(header) + b"".join(sectors)


def hwp_fixture(section_count=2, compressed=False, section_bodies=None, assets=()):
    def pad(data, tag, level=0):
        return data + record(tag, bytes(4096 - len(data) - 4), level)

    def encode(data):
        if not compressed:
            return data
        # Stored DEFLATE blocks keep these tiny fixtures above CFB's mini-stream cutoff.
        compressor = zlib.compressobj(level=0, wbits=-15)
        return compressor.compress(data) + compressor.flush()

    header = b"HWP Document File".ljust(32, b"\0") + bytes([0, 0, 0, 5]) + struct.pack("<I", int(compressed))
    mappings = record(16, struct.pack("<H", section_count) + bytes(24))
    mappings += record(17, struct.pack("<8I", len(assets), *([1] * 7)))
    for name, data in assets:
        identifier = int(name[3:7], 16)
        extension = name.split(".", 1)[1].encode("utf-16le")
        mappings += record(18, struct.pack("<HHH", 1, identifier, len(extension) // 2) + extension, 1)
    mappings += record(19, struct.pack("<BH", 0, 1) + b"A\0", 1) * 7
    shape = bytearray(68)
    struct.pack_into("<i", shape, 42, 1000)
    mappings += record(21, shape, 1) + record(25, struct.pack("<I6iH", 0, 0, 0, 0, 0, 0, 100, 0), 1)
    streams = [("FileHeader", header.ljust(4096, b"\0")), ("DocInfo", encode(pad(mappings, 27)))]
    for index in range(section_count):
        text = control_text(b"secd", 2) + f"Section {index}\r".encode("utf-16le")
        body = paragraph_records(text)
        body += record(71, b"dces" + bytes(20), 1)
        body += record(73, struct.pack("<10I", 61200, 79200, 7200 + index * 100,
                                      7200, 3600, 3600, 1800, 1800, 0, index % 2), 2)
        if section_bodies is not None:
            body = section_bodies[index]
        streams.append((f"Section{index}", encode(pad(body, 70, 1))))
    return compound_file(streams, [(name, encode(data.ljust(4096, b"\0"))) for name, data in assets])


class MultiSectionTests(unittest.TestCase):
    def test_complete_hwp_numeric_order_orientation_and_shared_styles(self):
        for compressed in (False, True):
            with self.subTest(compressed=compressed), tempfile.TemporaryDirectory() as directory:
                source, output = Path(directory) / "input.hwp", Path(directory) / "output.docx"
                source.write_bytes(hwp_fixture(11, compressed))
                root, warnings = read_hwp(source)
                self.assertEqual([section.get("source-stream") for section in root.findall("BodyText/SectionDef")],
                                 [f"Section{index}" for index in range(11)])
                Converter(root, warnings).convert().save(output)
                document = Document(output)
                self.assertEqual([p.text for p in document.paragraphs if p.text],
                                 [f"Section {index}" for index in range(11)])
                for index, section in enumerate(document.sections):
                    self.assertEqual(section.orientation, WD_ORIENT.LANDSCAPE if index % 2 else WD_ORIENT.PORTRAIT)
                    self.assertEqual(section.page_width.twips, 15840 if index % 2 else 12240)
                    self.assertEqual(section.left_margin.twips, 1440 + index * 20)
                    self.assertEqual(section.start_type, WD_SECTION_START.NEW_PAGE)
                self.assertEqual(validate_output(root, output)["sections"], 11)

    def test_rejects_missing_invalid_duplicate_and_nonconsecutive_sections(self):
        for names in ([], ["Section1"], ["Section0", "Section2"],
                      ["Section0", "Section0"], ["Section00"], ["Section0", "Unexpected"]):
            with self.subTest(names=names), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "input.hwp"
                source.touch()
                header = b"HWP Document File".ljust(32, b"\0") + bytes([0, 0, 0, 5]) + bytes(4)
                with patch("hwp_parser.olefile.OleFileIO") as container:
                    ole = container.return_value.__enter__.return_value
                    ole.get_size.return_value = len(header)
                    ole.openstream.return_value = BytesIO(header)
                    ole.listdir.return_value = [["BodyText", name] for name in names]
                    with self.assertRaisesRegex(ValueError, "section"):
                        HwpParser().parse(source)

    def test_declared_section_count_is_checked(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "input.hwp"
            source.write_bytes(hwp_fixture())
            # Preserve valid OLE framing, but change the DocInfo count from two to three.
            import olefile
            with olefile.OleFileIO(source, write_mode=True) as ole:
                data = bytearray(ole.openstream("DocInfo").read())
                struct.pack_into("<H", data, 4, 3)
                ole.write_stream("DocInfo", bytes(data))
            with self.assertRaisesRegex(ValueError, "section count"):
                HwpParser().parse(source)

    def test_page_inference_resets_at_section_boundaries(self):
        root, columns = sample_root()
        first = paragraph("First")
        first.find("LineSeg").set("y", "9000")
        columns.append(first)
        section = ET.SubElement(root.find("BodyText"), "SectionDef")
        columns = ET.SubElement(section, "ColumnSet")
        for text, y in (("Second", "100"), ("Third", "50")):
            node = paragraph(text)
            node.find("LineSeg").set("y", y)
            columns.append(node)
        document = Converter(root, [], preserve_source_pages=True).convert()
        paragraphs = {p.text: p for p in document.paragraphs if p.text}
        self.assertFalse(paragraphs["Second"].paragraph_format.page_break_before)
        self.assertTrue(paragraphs["Third"].paragraph_format.page_break_before)
        ET.SubElement(columns[0].find("LineSeg"), "ColumnsDef", {"count": "2"})
        with self.assertRaisesRegex(ValueError, "single-column"):
            Converter(root, [], preserve_source_pages=True).convert()

    def test_validation_detects_section_loss_and_changed_page_settings(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "input.hwp", Path(directory) / "output.docx"
            source.write_bytes(hwp_fixture())
            root, warnings = HwpParser().parse(source)
            for alteration in ("count", "margin", "orientation", "break"):
                with self.subTest(alteration=alteration):
                    document = Converter(root, warnings).convert()
                    if alteration == "count":
                        sect_pr = document.sections[0]._sectPr
                        sect_pr.getparent().remove(sect_pr)
                    elif alteration == "margin":
                        document.sections[1].left_margin = Inches(2)
                    elif alteration == "orientation":
                        document.sections[1].orientation = WD_ORIENT.PORTRAIT
                    else:
                        document.sections[1].start_type = WD_SECTION_START.CONTINUOUS
                    document.save(output)
                    with self.assertRaisesRegex(ValueError, "Section"):
                        validate_output(root, output)


class ResourceLimitTests(unittest.TestCase):
    def test_header_preflight_respects_budgets_before_opening_stream(self):
        for limits, option in ((HwpLimits(max_stream_mb=1), "max-stream-mb"),
                               (HwpLimits(max_expanded_mb=1), "max-expanded-mb")):
            with self.subTest(option=option), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "input.hwp"
                source.touch()
                with patch("convert_hwp.olefile.OleFileIO") as container:
                    ole = container.return_value.__enter__.return_value
                    ole.get_size.return_value = MIB + 1
                    with self.assertRaisesRegex(ValueError, option):
                        read_hwp(source, limits=limits)
                    ole.openstream.assert_not_called()

    def test_larger_than_50_mib_file_converts_with_explicit_budget(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "large.hwp", Path(directory) / "large.docx"
            source.write_bytes(hwp_fixture())
            # Valid compound streams plus sparse unused space exercise the actual file-size gate.
            with source.open("r+b") as stream:
                stream.truncate(50 * MIB + 1)
            with self.assertRaisesRegex(ValueError, "max-input-mb"):
                read_hwp(source)
            command = [sys.executable, "convert_hwp.py", str(source), str(output),
                       "--max-input-mb", "51", "--max-stream-mb", "2", "--max-expanded-mb", "3",
                       "--max-records", "100", "--reader-timeout", "120"]
            result = subprocess.run(command, capture_output=True, timeout=10)
            self.assertEqual(result.returncode, 0, result.stderr.decode())
            report = json.loads(output.with_suffix(".report.json").read_text())
            self.assertEqual(report["hwp5_limits"], {"max_input_mb": 51, "max_stream_mb": 2,
                                                     "max_expanded_mb": 3, "max_records": 100,
                                                     "size_unit": "MiB"})
            self.assertEqual(report["reader_timeout_seconds"], 120)
            self.assertEqual(report["sections"], 2)

    def test_worker_receives_record_budget_and_does_not_publish_on_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "input.hwp", Path(directory) / "output.docx"
            source.write_bytes(hwp_fixture())
            result = subprocess.run([sys.executable, "convert_hwp.py", str(source), str(output),
                                     "--max-records", "1"], capture_output=True, timeout=10)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn(b"record count", result.stderr)
            self.assertFalse(output.exists())
            self.assertFalse(output.with_suffix(".report.json").exists())

    def test_per_stream_limit_can_be_raised_and_remains_enforced(self):
        payload = b"a" * (MIB + 1)
        compressor = zlib.compressobj(wbits=-15)
        compressed = compressor.compress(payload) + compressor.flush()
        for raw, is_compressed in ((payload, False), (compressed, True)):
            for budget, succeeds in ((1, False), (2, True)):
                with self.subTest(compressed=is_compressed, budget=budget):
                    ole = Mock()
                    ole.get_size.return_value = len(raw)
                    ole.openstream.return_value = BytesIO(raw)
                    parser = HwpParser(limits=HwpLimits(max_stream_mb=budget))
                    if succeeds:
                        self.assertEqual(parser.stream(ole, "DocInfo", is_compressed), payload)
                    else:
                        with self.assertRaisesRegex(ValueError, "max-stream-mb"):
                            parser.stream(ole, "DocInfo", is_compressed)

    def test_cumulative_expanded_budget_checked_before_overallocation(self):
        parser = HwpParser(limits=HwpLimits(max_stream_mb=2, max_expanded_mb=1))
        raw = b"a" * (MIB // 2 + 1)
        compressor = zlib.compressobj(wbits=-15)
        compressed = compressor.compress(raw) + compressor.flush()
        ole = Mock()
        ole.get_size.return_value = len(compressed)
        ole.openstream.side_effect = lambda name: BytesIO(compressed)
        parser.stream(ole, "Section0", True)
        with self.assertRaisesRegex(ValueError, "max-expanded-mb"):
            parser.stream(ole, "Section1", True)
        self.assertEqual(parser.stream_bytes, len(raw))

    def test_record_tree_uses_selected_byte_and_record_budgets(self):
        data = record(70, b"a" * MIB, extended=True)
        with self.assertRaisesRegex(ValueError, "size limit"):
            record_tree(data, max_bytes=MIB)
        self.assertEqual(len(record_tree(data, max_bytes=2 * MIB)), 1)
        with self.assertRaisesRegex(ValueError, "record count"):
            record_tree(record(66) * 3, max_records=2)

    def test_timeout_is_forwarded_and_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "input.hwp"
            source.write_bytes(hwp_fixture())
            with patch("convert_hwp.subprocess.run", side_effect=subprocess.TimeoutExpired("reader", 3)) as worker:
                with self.assertRaisesRegex(ValueError, "exceeded 3 seconds"):
                    read_hwp(source, reader_timeout=3)
                self.assertEqual(worker.call_args.kwargs["timeout"], 3)

    def test_invalid_resource_arguments_fail_before_conversion(self):
        for option, value in (("--max-input-mb", "0"), ("--max-stream-mb", "-1"),
                              ("--max-expanded-mb", "1.5"), ("--max-records", "no"),
                              ("--reader-timeout", "0")):
            with self.subTest(option=option):
                result = subprocess.run([sys.executable, "convert_hwp.py", "missing.hwp", "unused.docx",
                                         option, value], capture_output=True, timeout=10)
                self.assertEqual(result.returncode, 2)
                self.assertIn(b"positive integer", result.stderr)
        for value in (0, -1, True, 1.5):
            with self.assertRaises(ValueError):
                HwpLimits(max_input_mb=value)


if __name__ == "__main__":
    unittest.main()
