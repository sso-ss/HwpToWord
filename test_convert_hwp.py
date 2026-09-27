from concurrent.futures import ThreadPoolExecutor
from contextlib import redirect_stderr
from io import StringIO
import json
import os
import tempfile
import struct
from threading import Barrier
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT

from convert_hwp import Converter, main, publish_output, read_hwp, validate_output
from hwp_reader import TextRecovery


def paragraph(text):
    node = ET.Element("Paragraph")
    line = ET.SubElement(node, "LineSeg")
    ET.SubElement(line, "Text", {"charshape-id": "0"}).text = text
    return node


def sample_root():
    root = ET.Element("HwpDoc")
    body = ET.SubElement(root, "BodyText")
    section = ET.SubElement(body, "SectionDef")
    columns = ET.SubElement(section, "ColumnSet")
    return root, columns


class ConversionTests(unittest.TestCase):
    def test_agreement_tabs_cell_padding_and_source_page_break(self):
        source = Path("output/5. 협력업체 업무협약서.hwp")
        if not source.exists():
            self.skipTest("Local fixture is unavailable.")
        root, warnings = read_hwp(source)
        document = Converter(root, warnings, preserve_source_pages=True).convert()
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        article = next(item for item in document.paragraphs if "제 9 조" in item.text)
        self.assertTrue(article.paragraph_format.page_break_before)
        clause = next(item for item in document.paragraphs if "협무협약" in item.text)
        self.assertGreater(len(clause.paragraph_format.tab_stops), 0)
        cell = document.tables[0].cell(0, 0)
        self.assertEqual(document.tables[0].alignment, WD_TABLE_ALIGNMENT.RIGHT)
        self.assertEqual(document.tables[1].alignment, WD_TABLE_ALIGNMENT.RIGHT)
        margin = cell._tc.tcPr.find(namespace + "tcMar/" + namespace + "left")
        self.assertEqual(margin.get(namespace + "w"), "28")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "agreement.docx"
            document.save(output)
            self.assertTrue(validate_output(root, output)["text_preserved"])

    def test_recovery_preserves_valid_surrogate_pairs_and_korean(self):
        recovery = TextRecovery()
        text = "한글 \U0001f600"
        self.assertEqual(recovery.decode(text.encode("utf-16le")), text)
        self.assertEqual(recovery.replacements, [])

    def test_recovery_records_every_invalid_sequence(self):
        recovery = TextRecovery()
        self.assertEqual(recovery.decode(b"\x80\xdbA\x00\x00\xdc"), "\ufffdA\ufffd")
        self.assertEqual([item["byte_offset_in_chunk"] for item in recovery.replacements], [0, 4])
        self.assertEqual(recovery.replacements[0]["invalid_bytes_hex"], "80db")
        self.assertEqual(recovery.decode(b"B"), "\ufffd")
        self.assertEqual(recovery.replacements[-1]["chunk_index"], 2)

    def test_text_order_and_native_merged_table(self):
        root, columns = sample_root()
        columns.append(paragraph("Before"))
        host = ET.SubElement(columns, "Paragraph")
        line = ET.SubElement(host, "LineSeg")
        table = ET.SubElement(line, "TableControl", {"width": "14400"})
        body = ET.SubElement(table, "TableBody", {"rows": "2", "cols": "2"})
        for row_index in range(2):
            row = ET.SubElement(body, "TableRow")
            count = 1 if row_index == 0 else 2
            for col_index in range(count):
                cell = ET.SubElement(row, "TableCell", {
                    "row": str(row_index), "col": str(col_index),
                    "colspan": "2" if row_index == 0 else "1", "rowspan": "1",
                    "width": "14400" if row_index == 0 else "7200",
                })
                cell.append(paragraph("Merged" if row_index == 0 else "Cell"))
        columns.append(paragraph("After"))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "test.docx"
            Converter(root, []).convert().save(output)
            report = validate_output(root, output)
            self.assertTrue(report["text_preserved"])
            document = Document(output)
            self.assertEqual(document.tables[0].cell(0, 0).text, "Merged")
            self.assertEqual(document.tables[0].cell(0, 0)._tc.grid_span, 2)

    def test_validation_detects_missing_text(self):
        root, columns = sample_root()
        columns.append(paragraph("Must remain"))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "empty.docx"
            Document().save(output)
            with self.assertRaisesRegex(ValueError, "Text preservation"):
                validate_output(root, output)

    def test_nested_table_uses_cell_api(self):
        root, columns = sample_root()
        host = paragraph("")
        columns.append(host)
        for depth in range(2):
            control = ET.SubElement(host.find("LineSeg"), "TableControl")
            body = ET.SubElement(control, "TableBody", {"rows": "1", "cols": "1"})
            row = ET.SubElement(body, "TableRow")
            cell = ET.SubElement(row, "TableCell", {"row": "0", "col": "0"})
            host = paragraph("Nested" if depth else "")
            cell.append(host)
        document = Converter(root, []).convert()
        self.assertEqual(document.tables[0].cell(0, 0).tables[0].cell(0, 0).text, "Nested")

    def test_textbox_and_hyperlink_labels_remain_in_order(self):
        root, columns = sample_root()
        host = paragraph("Before")
        columns.append(host)
        line = host.find("LineSeg")
        shape = ET.SubElement(line, "GShapeObjectControl")
        textbox = ET.SubElement(shape, "TextboxParagraphList")
        textbox.append(paragraph("Box"))
        link = ET.SubElement(line, "FieldHyperLink")
        ET.SubElement(link, "Text").text = "Link"
        ET.SubElement(line, "Text").text = "After"
        converter = Converter(root, [])
        document = converter.convert()
        self.assertEqual(document.paragraphs[0].text, "BeforeBoxLinkAfter")
        self.assertTrue(any("Text boxes" in warning for warning in converter.warnings))

    def test_hanging_indent_keeps_first_line_at_left_margin(self):
        root, columns = sample_root()
        info = ET.SubElement(root, "DocInfo")
        ET.SubElement(info, "ParaShape", {"indent": "-5400", "doubled-margin-left": "0", "linespacing-type": "ratio"})
        columns.append(paragraph("Indented text"))
        converted = Converter(root, []).convert().paragraphs[0].paragraph_format
        self.assertEqual(converted.left_indent.pt, 54)
        self.assertEqual(converted.first_line_indent.pt, -54)

    def test_table_page_break_precedes_table_and_is_not_repeated(self):
        for mode in ("inferred", "explicit", "style"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as directory:
                root, columns = sample_root()
                before = paragraph("Previous page")
                before.find("LineSeg").set("y", "1000")
                columns.append(before)
                host = ET.SubElement(columns, "Paragraph")
                if mode == "explicit":
                    host.set("new-page", "1")
                elif mode == "style":
                    info = ET.SubElement(root, "DocInfo")
                    ET.SubElement(info, "ParaShape")
                    ET.SubElement(info, "ParaShape", {"start-new-page": "1"})
                    host.set("parashape-id", "1")
                line = ET.SubElement(host, "LineSeg", {"y": "0"})
                for text in ("First table", "Second table"):
                    table = ET.SubElement(line, "TableControl")
                    body = ET.SubElement(table, "TableBody", {"rows": "1", "cols": "1"})
                    row = ET.SubElement(body, "TableRow")
                    cell = ET.SubElement(row, "TableCell", {"row": "0", "col": "0"})
                    cell.append(paragraph(text))
                ET.SubElement(line, "Text").text = "After tables"
                output = Path(directory) / "page-break.docx"
                Converter(root, [], preserve_source_pages=mode == "inferred").convert().save(output)
                document = Document(output)
                self.assertTrue(document.paragraphs[1].paragraph_format.page_break_before)
                self.assertTrue(document.paragraphs[1].paragraph_format.keep_with_next)
                self.assertIs(document.tables[0]._tbl.getprevious(), document.paragraphs[1]._p)
                self.assertEqual(sum(p.paragraph_format.page_break_before is True
                                     for p in document.paragraphs), 1)
                self.assertEqual(document.paragraphs[-1].text, "After tables")
                self.assertFalse(document.paragraphs[-1].paragraph_format.page_break_before)
                self.assertTrue(validate_output(root, output)["text_preserved"])

    def test_validation_rejects_missing_textbox_separators(self):
        root, columns = sample_root()
        host = paragraph("")
        columns.append(host)
        drawing = ET.SubElement(host.find("LineSeg"), "GShapeObjectControl")
        textbox = ET.SubElement(drawing, "TextboxParagraphList")
        first = paragraph("First")
        line = first.find("LineSeg")
        for control, text in (("TAB", "Second"), ("LINE_BREAK", "Third"), ("FIXWIDTH_SPACE", "Fourth")):
            ET.SubElement(line, "ControlChar", {"name": control})
            ET.SubElement(line, "Text").text = text
        textbox.append(first)
        textbox.append(paragraph("Fifth"))
        expected = "First\tSecond\nThird\u2007Fourth\nFifth"
        for separator in ("\t", "\n", "\u2007", "paragraph"):
            with self.subTest(separator=separator), tempfile.TemporaryDirectory() as directory:
                document = Document()
                # Same letters, with exactly one separator lost.
                damaged = expected.rsplit("\n", 1) if separator == "paragraph" else None
                document.add_paragraph("".join(damaged) if damaged else expected.replace(separator, "", 1))
                output = Path(directory) / "damaged.docx"
                document.save(output)
                with self.assertRaisesRegex(ValueError, "Text preservation"):
                    validate_output(root, output)

    def test_validation_ignores_tab_stop_definitions(self):
        root, columns = sample_root()
        info = ET.SubElement(root, "DocInfo")
        ET.SubElement(info, "ParaShape", {"tabdef-id": "0"})
        tabs = ET.SubElement(ET.SubElement(info, "TabDef"), "Array")
        ET.SubElement(tabs, "Tab", {"pos": "7200", "kind": "left"})
        host = paragraph("First")
        ET.SubElement(host.find("LineSeg"), "ControlChar", {"name": "TAB"})
        ET.SubElement(host.find("LineSeg"), "Text").text = "Second"
        columns.append(host)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "tab-stops.docx"
            Converter(root, []).convert().save(output)
            self.assertEqual(Document(output).paragraphs[0].text, "First\tSecond")
            self.assertTrue(validate_output(root, output)["text_preserved"])

    def test_unsupported_format(self):
        with self.assertRaisesRegex(ValueError, "not HWPX"):
            read_hwp(Path("input.hwpx"))

    def test_protected_documents_are_rejected_before_parsing(self):
        for flag in (1 << 1, 1 << 2, 1 << 4, 1 << 8, 1 << 10):
            with self.subTest(flag=flag), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "protected.hwp"
                source.touch()
                header = b"HWP Document File".ljust(32, b"\0") + bytes([0, 0, 0, 5]) + struct.pack("<I", flag)
                with patch("convert_hwp.olefile.OleFileIO") as reader, patch("convert_hwp.subprocess.run") as process:
                    reader.return_value.__enter__.return_value.get_size.return_value = len(header)
                    reader.return_value.__enter__.return_value.openstream.return_value.read.return_value = header
                    with self.assertRaisesRegex(ValueError, "protected"):
                        read_hwp(source)
                    process.assert_not_called()

    def test_sample_conversion(self):
        source = Path("output/2025 멘토링 프로그램_하반기.hwp")
        if not source.exists():
            self.skipTest("Local private fixture is not included in the repository.")
        root, warnings = read_hwp(source)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "sample.docx"
            Converter(root, warnings).convert().save(output)
            report = validate_output(root, output)
            self.assertEqual(report["tables"], 2)
            self.assertEqual(report["text_characters"], 740)
            document = Document(output)
            self.assertEqual(len(document.tables[1].rows), 10)
            self.assertEqual(len(document.tables[1].columns), 3)
            for table in document.tables:
                following = table._tbl.getnext()
                self.assertEqual(following.tag.rsplit("}", 1)[-1], "p")
                spacing = following.find("{*}pPr/{*}spacing")
                self.assertEqual(spacing.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line"), "20")
            title = document.tables[0].cell(0, 0).paragraphs[0].runs[0]
            self.assertEqual(title.font.size.pt, 20)
            self.assertTrue(title.bold)
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            spacing = title._element.find(namespace + "rPr/" + namespace + "spacing")
            self.assertEqual(spacing.get(namespace + "val"), "-12")
            width = document.tables[0]._tbl.tblPr.find(namespace + "tblW")
            self.assertEqual(width.get(namespace + "type"), "dxa")
            self.assertEqual(width.get(namespace + "w"), "9530")
            body = next(item for item in document.paragraphs if "개최" in item.text)
            self.assertAlmostEqual(body.paragraph_format.line_spacing.pt, 22.4, places=1)
            self.assertEqual(document.tables[1].cell(0, 0).paragraphs[0].paragraph_format.line_spacing.pt, 12)

    def test_font_substitution_is_explicit_and_reported(self):
        source = Path("output/2025 멘토링 프로그램_하반기.hwp")
        if not source.exists():
            self.skipTest("Local fixture is unavailable.")
        root, warnings = read_hwp(source)
        converter = Converter(root, warnings, {"HY헤드라인M": "HeadLineA"})
        document = converter.convert()
        title = document.tables[0].cell(0, 0).paragraphs[0].runs[0]
        self.assertEqual(title.font.name, "HeadLineA")
        self.assertIn("Font substitution: HY헤드라인M -> HeadLineA", converter.warnings)


class PublicationTests(unittest.TestCase):
    def test_files_created_during_conversion_are_preserved(self):
        for destinations in (("document",), ("report",), ("document", "report")):
            with self.subTest(destinations=destinations), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / "result.docx"
                report = output.with_suffix(".report.json")
                paths = {"document": output, "report": report}
                root, columns = sample_root()
                columns.append(paragraph("New document"))

                def competing_conversion(*args, **kwargs):
                    for name in destinations:
                        paths[name].write_bytes(b"Already published by another conversion")
                    return root, []

                with patch("sys.argv", ["convert_hwp.py", "input.hwp", str(output)]), \
                        patch("convert_hwp.read_hwp", side_effect=competing_conversion), \
                        redirect_stderr(StringIO()) as errors, self.assertRaises(SystemExit) as failure:
                    main()
                self.assertEqual(failure.exception.code, 1)
                self.assertIn("already exists", errors.getvalue())
                for name, path in paths.items():
                    if name in destinations:
                        self.assertEqual(path.read_bytes(), b"Already published by another conversion")
                    else:
                        self.assertFalse(path.exists())

    def test_concurrent_publication_keeps_document_and_report_together(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "result.docx"
            ready = Barrier(2)
            candidates = []
            for label in ("first", "second"):
                candidate = Path(directory) / (label + ".docx")
                document = Document()
                document.add_paragraph(label)
                document.save(candidate)
                candidates.append(candidate)

            def publish(candidate):
                ready.wait(timeout=5)
                try:
                    publish_output(candidate, output, {"label": candidate.stem})
                    return candidate.stem
                except ValueError:
                    return None

            with ThreadPoolExecutor(max_workers=2) as workers:
                results = list(workers.map(publish, candidates))
            winners = [result for result in results if result is not None]
            self.assertEqual(len(winners), 1)
            self.assertEqual(Document(output).paragraphs[0].text, winners[0])
            self.assertEqual(json.loads(output.with_suffix(".report.json").read_text()), {"label": winners[0]})

    def test_failed_document_publication_removes_only_its_own_report(self):
        for replace_report in (False, True):
            with self.subTest(replace_report=replace_report), tempfile.TemporaryDirectory() as directory:
                candidate = Path(directory) / "candidate.docx"
                candidate.write_bytes(b"staged document")
                output = Path(directory) / "result.docx"
                report = output.with_suffix(".report.json")
                link = os.link

                def fail_document_link(source, destination):
                    if destination == output:
                        if replace_report:
                            report.unlink()
                            report.write_text("An unrelated replacement report")
                        raise OSError("Document publication failed")
                    link(source, destination)

                with patch("convert_hwp.os.link", side_effect=fail_document_link), \
                        self.assertRaisesRegex(OSError, "Document publication failed"):
                    publish_output(candidate, output, {})
                self.assertFalse(output.exists())
                if replace_report:
                    self.assertEqual(report.read_text(), "An unrelated replacement report")
                else:
                    self.assertFalse(report.exists())

    def test_dangling_output_symlinks_are_not_replaced(self):
        for suffix in (".docx", ".report.json"):
            with self.subTest(suffix=suffix), tempfile.TemporaryDirectory() as directory:
                candidate = Path(directory) / "candidate.docx"
                candidate.write_bytes(b"staged document")
                output = Path(directory) / "result.docx"
                occupied = output.with_suffix(suffix)
                missing = Path(directory) / "missing"
                occupied.symlink_to(missing)
                with self.assertRaisesRegex(ValueError, "already exists"):
                    publish_output(candidate, output, {})
                self.assertTrue(occupied.is_symlink())
                self.assertFalse(missing.exists())
                other = output.with_suffix(".report.json") if suffix == ".docx" else output
                self.assertFalse(other.exists())


if __name__ == "__main__":
    unittest.main()
