"""Synthetic HWP 3 fixtures from the published field offsets, not real files.

These tests establish our supported interpretation and corruption handling.
They do not establish compatibility with an original Hancom 3.x writer.
"""

import gzip
import json
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zlib

from docx import Document

from convert_hwp import Converter, read_hwp, source_text, validate_output
from hwp3_parser import Hwp3Parser, decode_hchar, decompress_stream
from hwp_reader import parse_source


def put(data, offset, value, kind="H"):
    struct.pack_into("<" + kind, data, offset, value)


def char_shape(flags=0, size=250):
    data = bytearray(31)
    put(data, 0, size)
    data[9:16] = bytes([100] * 7)
    data[26] = flags
    return bytes(data)


def para_shape():
    data = bytearray(187)
    put(data, 6, 160)
    data[10] = 100
    data[172] = 1
    return bytes(data)


def hchars(text):
    # kchar uses the stdlib codec; modern syllables share the Johab bit layout.
    return b"".join(struct.pack("<H", ord(c) if ord(c) < 128 else int.from_bytes(c.encode("johab"), "big")) for c in text)


END = b"\x01" + bytes(42)


def paragraph(text="", raw=None, count=None, shape=None, flags=0, inherit=False,
              char=None, changes=None, lines=()):
    raw = hchars(text + "\r") if raw is None else raw
    count = len(raw) // 2 if count is None else count
    header = bytearray(43)
    header[0] = int(inherit)
    put(header, 1, count)
    put(header, 3, len(lines))
    header[5] = int(changes is not None)
    header[6] = flags
    header[12:43] = char or char_shape()
    definitions = b"" if inherit else (shape or para_shape())
    runs = b"" if changes is None else b"".join(b"\x01" if change is None else b"\0" + change for change in changes)
    return bytes(header) + definitions + b"".join(lines) + runs + raw


def table_cell(x=0, y=0, width=1800, height=900, borders=(1, 2, 3, 4), diagonal=0):
    cell = bytearray(27)
    for offset, value in zip((4, 6, 8, 10), (x, y, width, height)):
        put(cell, offset, value)
    cell[18] = 1
    cell[19] = 1
    cell[20:24] = bytes(borders)
    cell[25] = diagonal
    return bytes(cell)


def table(cells, kind=0, caption=b"", anchor=0):
    data = bytearray(84)
    data[8] = anchor
    put(data, 16, 10)
    for offset in (34, 36, 38, 40):
        put(data, offset, 25)
    put(data, 42, 3600)
    put(data, 44, 1800)
    put(data, 78, kind)
    put(data, 80, len(cells))
    return (struct.pack("<HIH", 10, 0, 10) + bytes(data)
            + b"".join(cell for cell, text in cells)
            + b"".join(text + END for cell, text in cells) + caption + END)


def fixture(paragraphs=None, compression=None, body_tail=b"", preview=b"", styles=b""):
    info = bytearray(128)
    info[4] = 3
    info[125] = 1
    put(info, 6, 21047)
    put(info, 8, 14882)
    for offset in (10, 12, 14, 16):
        put(info, offset, 900)
    put(info, 18, 450)
    put(info, 20, 450)
    info[124] = int(compression is not None)
    fonts = (b"\x01\0" + "바탕".encode("johab").ljust(40, b"\0")) * 7
    body = fonts + struct.pack("<H", len(styles) // 238) + styles
    body += (paragraph("한글 ABC") if paragraphs is None else paragraphs) + END + body_tail
    if compression == "gzip":
        body = gzip.compress(body, mtime=0)
    elif compression in ("raw", "footer"):
        encoder = zlib.compressobj(wbits=-15)
        footer = struct.pack("<II", zlib.crc32(body), len(body))
        body = encoder.compress(body) + encoder.flush()
        if compression == "footer":
            body += footer
    return b"HWP Document File V3.00 \x1a\x01\x02\x03\x04\x05" + bytes(info) + bytes(1008) + body + preview


class Hwp3Tests(unittest.TestCase):
    def parse(self, data):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "synthetic.hwp"
            source.write_bytes(data)
            return parse_source(source)

    def save(self, data):
        root, warnings = self.parse(data)
        output = Converter(root, warnings).convert()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.docx"
            output.save(path)
            report = validate_output(root, path)
            return Document(path), report

    def test_uncompressed_gzip_raw_and_footer_with_preview(self):
        preview = struct.pack("<II", 0x80000002, 4) + hchars("OK") + struct.pack("<II", 0x80000000, 20)
        for compression in (None, "gzip", "raw", "footer"):
            for tail in (b"", bytes(8)):
                with self.subTest(compression=compression, tail=tail):
                    root, warnings = self.parse(fixture(compression=compression, body_tail=tail, preview=preview))
                    self.assertEqual(root.get("parser"), "native-hwp3")
                    self.assertEqual(source_text(root.find("BodyText")), "한글 ABC")
                    self.assertIn("unverified", " ".join(warnings))
                    self.assertEqual(root.find("DocInfo/FaceName").get("name"), "바탕")

    def test_modern_syllables_roundtrip_against_stdlib_codec(self):
        self.assertEqual(decode_hchar(0x8861), "가")
        for codepoint in range(0xac00, 0xd7a4):
            char = chr(codepoint)
            self.assertEqual(decode_hchar(int.from_bytes(char.encode("johab"), "big")), char)
        for code in (0x343b, 0x4000, 0xd931, 0x8000, 0xffff, 0x7f):
            with self.subTest(code=code), self.assertRaisesRegex(ValueError, "Unsupported HWP 3 hchar"):
                decode_hchar(code)

    def test_styles_units_inheritance_and_explicit_page_break(self):
        shape = bytearray(para_shape())
        put(shape, 0, 100)
        put(shape, 4, -50, "h")
        put(shape, 8, 25)
        put(shape, 183, 50)
        struct.pack_into("<BBH", shape, 12, 2, 1, 1800)
        content = paragraph("가A", shape=bytes(shape), changes=[char_shape(3), char_shape(4), None])
        content += paragraph("둘", inherit=True, flags=2, char=char_shape(32))
        document, report = self.save(fixture(content))
        self.assertTrue(report["text_preserved"])
        first, second = document.paragraphs
        self.assertEqual(first.text, "가A")
        self.assertTrue(first.runs[0].bold)
        self.assertTrue(first.runs[0].italic)
        self.assertTrue(first.runs[1].underline)
        self.assertEqual(first.runs[0].font.size.pt, 10)
        self.assertEqual(first.paragraph_format.line_spacing, 1.6)
        self.assertEqual(first.paragraph_format.left_indent.pt, 6)
        self.assertEqual(first.paragraph_format.first_line_indent.pt, -2)
        self.assertEqual(first.paragraph_format.space_before.pt, 2)
        self.assertEqual(first.paragraph_format.tab_stops[0].position.inches, 1)
        self.assertTrue(second.paragraph_format.page_break_before)
        self.assertTrue(second.runs[0].font.superscript)
        self.assertAlmostEqual(document.sections[0].page_width.inches, 14882 / 1800, places=3)

    def test_absolute_spacing_and_no_invented_page_inference(self):
        shape = bytearray(para_shape())
        put(shape, 6, 0x8000 | 300)
        root, warnings = self.parse(fixture(paragraph("Line", shape=bytes(shape))))
        converter = Converter(root, warnings, preserve_source_pages=True)
        doc = converter.convert()
        self.assertEqual(doc.paragraphs[0].paragraph_format.line_spacing.pt, 12)
        self.assertIn("only available for HWP 5", " ".join(converter.warnings))

    def test_korean_and_english_character_spacing_use_their_own_settings(self):
        shape = bytearray(char_shape())
        shape[9:11] = bytes([80, 120])
        shape[16:18] = bytes([0, 10])
        root, warnings = self.parse(fixture(paragraph("가A", char=bytes(shape))))
        runs = list(root.findall(".//Text"))
        self.assertEqual([run.get("lang") for run in runs], ["ko", "en"])
        doc = Converter(root, warnings).convert()
        self.assertEqual(doc.paragraphs[0].runs[0]._element.xpath("./w:rPr/w:w/@w:val"), ["80"])
        self.assertEqual(doc.paragraphs[0].runs[1]._element.xpath("./w:rPr/w:w/@w:val"), ["120"])
        self.assertEqual(doc.paragraphs[0].runs[1]._element.xpath("./w:rPr/w:spacing/@w:val"), ["20"])

    def test_tabs_and_distinct_spaces_survive_docx(self):
        raw = (hchars("A") + struct.pack("<4H", 9, 900, 1, 9) + hchars("B")
               + struct.pack("<2H", 30, 30) + hchars("C") + struct.pack("<2H", 31, 31) + hchars("D\r"))
        document, report = self.save(fixture(paragraph(raw=raw)))
        self.assertEqual(document.paragraphs[0].text, "A\tB\u00a0C\u2007D")
        self.assertEqual(report["fixed_width_spaces_approximated"], 1)

    def test_merged_nested_tables_and_cell_geometry(self):
        nested = table([(table_cell(), paragraph("Inner"))])
        cells = [(table_cell(0, 0, 3600, 900), paragraph("Merged")),
                 (table_cell(0, 900), paragraph("Left")),
                 (table_cell(1800, 900), paragraph(raw=nested + hchars("\r"), count=5))]
        data = table(cells)
        document, report = self.save(fixture(paragraph(raw=data + hchars("\r"), count=5, flags=2)))
        self.assertEqual(report["tables"], 2)
        self.assertEqual(document.tables[0].cell(0, 0).text, "Merged")
        self.assertEqual(document.tables[0].cell(1, 0).text, "Left")
        self.assertEqual(document.tables[0].cell(1, 1).tables[0].cell(0, 0).text, "Inner")
        self.assertTrue(document.paragraphs[0].paragraph_format.page_break_before)
        self.assertFalse(document.paragraphs[-1].paragraph_format.page_break_before)
        self.assertEqual(document.tables[0].cell(1, 0).width.inches, 1)

    def test_textbox_paragraph_boundaries_tabs_and_empty_paragraph(self):
        contents = paragraph(raw=hchars("A") + struct.pack("<4H", 9, 500, 0, 9) + hchars("B\r"))
        contents += paragraph("") + paragraph("C")
        data = table([(table_cell(), contents)], kind=1)
        document, report = self.save(fixture(paragraph(raw=data + hchars("\r"), count=5)))
        self.assertEqual(document.paragraphs[0].text, "A\tB\n\nC")
        self.assertTrue(report["text_preserved"])

    def test_unsupported_features_fail_instead_of_disappearing(self):
        for code in (5, 6, 7, 8, 11, 14, 16, 17, 18, 24, 28, 29):
            with self.subTest(code=code), self.assertRaisesRegex(ValueError, "Unsupported HWP 3 control"):
                self.parse(fixture(paragraph(raw=struct.pack("<2H", code, 13))))
        for flags in (1, 4):
            with self.assertRaisesRegex(ValueError, "column or section"):
                self.parse(fixture(paragraph(flags=flags)))
        for kind in (2, 3):
            with self.assertRaisesRegex(ValueError, "equations, buttons"):
                self.parse(fixture(paragraph(raw=table([(table_cell(), paragraph())], kind=kind) + hchars("\r"), count=5)))
        for code in (0x343b, 0x4000):
            with self.assertRaisesRegex(ValueError, "legacy Hanja"):
                self.parse(fixture(paragraph(raw=struct.pack("<2H", code, 13))))

    def test_rejects_caption_diagonal_overlap_gap_and_huge_grid(self):
        cases = [table([(table_cell(), paragraph())], caption=paragraph("Caption")),
                 table([(table_cell(diagonal=16), paragraph())]),
                 table([(table_cell(), paragraph()), (table_cell(), paragraph())]),
                 table([(table_cell(), paragraph()), (table_cell(3600, 0), paragraph())]),
                 table([(table_cell(0, 0, 0, 0), paragraph())])]
        for data in cases:
            with self.subTest(size=len(data)), self.assertRaises(ValueError):
                self.parse(fixture(paragraph(raw=data + hchars("\r"), count=5)))
        # 102 diagonal rectangles create a 203 x 203 grid despite few cells.
        data = table([(table_cell(i * 10, i * 10, 5, 5), paragraph()) for i in range(102)])
        with self.assertRaisesRegex(ValueError, "grid exceeds"):
            self.parse(fixture(paragraph(raw=data + hchars("\r"), count=5)))

    def test_truncation_and_unsupported_signatures(self):
        data = fixture()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.hwp"
            for end in (0, 29, 100, 158, 1165, 1170, 1450, len(data) - 1):
                path.write_bytes(data[:end])
                with self.subTest(end=end), self.assertRaises(ValueError):
                    parse_source(path)
            path.write_bytes(b"PK\x03\x04" + bytes(60))
            with self.assertRaisesRegex(ValueError, "HWPX"):
                parse_source(path)

    def test_malformed_lengths_inheritance_and_control_markers(self):
        cases = [paragraph(inherit=True), paragraph(raw=hchars("AB")),
                 paragraph(raw=hchars("\rA")), paragraph("A", changes=[None, None]),
                 paragraph(raw=struct.pack("<4H", 9, 0, 0, 8)),
                 paragraph(raw=struct.pack("<3H", 30, 31, 13))]
        for data in cases:
            with self.subTest(data=data[:15]), self.assertRaises(ValueError):
                self.parse(fixture(data))

    def test_protection_is_rejected_before_worker_and_directly(self):
        for offset, kind in ((24, "I"), (96, "H")):
            data = bytearray(fixture(compression="gzip"))
            put(data, 30 + offset, 1, kind)
            with self.assertRaisesRegex(ValueError, "protected"):
                self.parse(bytes(data))
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "protected.hwp"
                path.write_bytes(data)
                with patch("convert_hwp.subprocess.run") as worker, patch("convert_hwp.olefile.OleFileIO") as ole:
                    with self.assertRaisesRegex(ValueError, "protected"):
                        read_hwp(path)
                    worker.assert_not_called()
                    ole.assert_not_called()

    def test_compression_limits_checksums_and_trailing_data(self):
        data = gzip.compress(b"large" * 100)
        with self.assertRaisesRegex(ValueError, "size limit"):
            decompress_stream(data, limit=10)
        for broken in (data[:-1], data[:-8] + bytes(8)):
            with self.assertRaises(ValueError):
                decompress_stream(broken)
        for compression in (None, "gzip", "raw", "footer"):
            with self.subTest(compression=compression), self.assertRaises(ValueError):
                self.parse(fixture(compression=compression, preview=b"garbage"))
        data = bytearray(fixture(compression="footer"))
        data[-1] ^= 1
        with self.assertRaisesRegex(ValueError, "checksum"):
            self.parse(bytes(data))

    def test_bad_font_line_and_additional_block_fields(self):
        shape = bytearray(char_shape())
        shape[2] = 1
        with self.assertRaisesRegex(ValueError, "font reference"):
            self.parse(fixture(paragraph(char=bytes(shape))))
        line = bytearray(14)
        put(line, 0, 1)
        with self.assertRaisesRegex(ValueError, "line positions"):
            self.parse(fixture(paragraph("A", lines=[bytes(line)])))
        for tail in (struct.pack("<II", 1, 100), struct.pack("<II", 1, 0), struct.pack("<II", 0, 1)):
            with self.subTest(tail=tail), self.assertRaises(ValueError):
                self.parse(fixture(body_tail=tail))
        preview = struct.pack("<II", 0x80000002, 2) + hchars("A")
        with self.assertRaisesRegex(ValueError, "terminator"):
            self.parse(fixture(preview=preview))

    def test_recovery_does_not_hide_unsupported_hwp3_text_or_publish_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "symbol.hwp"
            output = Path(directory) / "result.docx"
            source.write_bytes(fixture(paragraph(raw=struct.pack("<2H", 0x343b, 13))))
            result = subprocess.run([sys.executable, "convert_hwp.py", str(source), str(output), "--recover-text"],
                                    capture_output=True, timeout=10)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn(b"0x343B", result.stderr)
            self.assertFalse(output.exists())
            self.assertFalse(output.with_suffix(".report.json").exists())

    def test_character_and_depth_budgets(self):
        with patch("hwp3_parser.MAX_CHARACTERS", 1), self.assertRaisesRegex(ValueError, "count exceeds"):
            self.parse(fixture())
        data = table([(table_cell(), paragraph("Inner"))])
        with patch("hwp3_parser.MAX_DEPTH", 0), self.assertRaisesRegex(ValueError, "nesting"):
            self.parse(fixture(paragraph(raw=data + hchars("\r"), count=5)))

    def test_stdlib_only_worker_and_converter_cli_report(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "한글 3.hwp"
            output = Path(directory) / "result.docx"
            source.write_bytes(fixture(compression="gzip"))
            # -S disables site-packages: olefile/python-docx are unavailable.
            worker = subprocess.run([sys.executable, "-S", "hwp_reader.py", str(source)], capture_output=True, timeout=10)
            self.assertEqual(worker.returncode, 0, worker.stderr)
            self.assertEqual(ET.fromstring(worker.stdout).get("parser"), "native-hwp3")
            command = [sys.executable, "convert_hwp.py", str(source), str(output)]
            result = subprocess.run(command, capture_output=True, timeout=10)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(output.with_suffix(".report.json").read_text())
            self.assertEqual(report["parser"], "native-hwp3")
            self.assertTrue(report["text_preserved"])
            self.assertEqual(report["text_replacements"], [])
            self.assertEqual(Document(output).paragraphs[0].text, "한글 ABC")
            before = output.read_bytes()
            again = subprocess.run(command, capture_output=True, timeout=10)
            self.assertNotEqual(again.returncode, 0)
            self.assertEqual(before, output.read_bytes())


if __name__ == "__main__":
    unittest.main()
