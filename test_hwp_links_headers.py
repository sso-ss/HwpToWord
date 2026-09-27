from contextlib import redirect_stdout
from copy import deepcopy
from io import StringIO
import base64
import json
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

from convert_hwp import Converter, main, read_hwp, validate_output
from hwp_parser import record_tree
from test_convert_hwp import paragraph, sample_root
from test_hwp_parser import control_text, paragraph_records, parser_with_styles, record, table_records
from test_hwp_limits_sections import hwp_fixture


PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a1XkAAAAASUVORK5CYII=")


def link_paragraph(command="https://example.org/a\\;b?x=1&y=2;1;0;0", label="Linked 한글", level=0):
    text = (b"B\0" + control_text(b"%hlk", 3) + label.encode("utf-16le")
            + control_text(bytes(4), 4) + b"A\0\r\0")
    encoded = command.encode("utf-16le")
    return (paragraph_records(text, level)
            + record(71, b"klh%" + struct.pack("<IBH", 0, 0, len(encoded) // 2) + encoded
                     + struct.pack("<I", 7), level + 1))


def header_records(kind=b"head", text="Header", selection=0, contents=None):
    contents = paragraph_records((text + "\r").encode("utf-16le"), 2) if contents is None else contents
    return (record(71, kind[::-1] + struct.pack("<I", selection), 1)
            + record(72, struct.pack("<II", 1, 0), 2) + contents)


def picture_records(level=1):
    return (record(71, struct.pack("<4s5I", b" osg", 1, 0, 0, 7200, 3600), level)
            + record(76, b"cip$", level + 1)
            + record(85, bytes(71) + struct.pack("<H", 1) + bytes(5), level + 2))


class HyperlinkTests(unittest.TestCase):
    def test_binary_fields_create_active_links_and_keep_labels(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "links.hwp", Path(directory) / "links.docx"
            source.write_bytes(hwp_fixture(1, section_bodies=[link_paragraph()]))
            root, warnings = read_hwp(source)
            self.assertEqual(root.find(".//FieldHyperLink").get("target"), "https://example.org/a;b?x=1&y=2")
            Converter(root, warnings).convert().save(output)
            document = Document(output)
            self.assertEqual(document.paragraphs[0].text, "BLinked 한글A")
            link = document.paragraphs[0].hyperlinks[0]
            self.assertEqual(link.address, "https://example.org/a;b?x=1&y=2")
            self.assertEqual(validate_output(root, output)["hyperlinks"], 1)

    def test_styled_labels_tabs_and_line_breaks_survive(self):
        root, columns = sample_root()
        host = paragraph("Before")
        link = ET.SubElement(host.find("LineSeg"), "FieldHyperLink", {"target": "mailto:reader@example.org"})
        ET.SubElement(link, "Text").text = "One"
        ET.SubElement(link, "ControlChar", {"name": "TAB"})
        ET.SubElement(link, "Text").text = "Two"
        ET.SubElement(link, "ControlChar", {"name": "LINE_BREAK"})
        ET.SubElement(link, "Text").text = "Three"
        info = ET.SubElement(root, "DocInfo")
        ET.SubElement(info, "CharShape", {"bold": "1", "basesize": "1400"})
        columns.append(host)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "styled.docx"
            Converter(root, []).convert().save(output)
            document = Document(output)
            self.assertEqual(document.paragraphs[0].text, "BeforeOne\tTwo\nThree")
            self.assertTrue(document.paragraphs[0].hyperlinks[0].runs[0].bold)
            self.assertTrue(validate_output(root, output)["hyperlink_destinations_preserved"])

    def test_unsupported_destinations_remain_labels(self):
        for target in ("#bookmark", "file:///document.hwp", "javascript:alert(1)", "", "http://[broken"):
            with self.subTest(target=target):
                root, columns = sample_root()
                host = paragraph("Before")
                link = ET.SubElement(host.find("LineSeg"), "FieldHyperLink", {"target": target})
                ET.SubElement(link, "Text").text = "Label"
                columns.append(host)
                converter = Converter(root, [])
                document = converter.convert()
                self.assertEqual(document.paragraphs[0].text, "BeforeLabel")
                self.assertEqual(document.paragraphs[0].hyperlinks, [])
                self.assertTrue(any("target" in warning for warning in converter.warnings))

    def test_validation_rejects_lost_or_changed_hyperlinks(self):
        parser = parser_with_styles()
        parser.section.find("ColumnSet").append(parser.paragraph(record_tree(link_paragraph())[0]))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "changed.docx"
            for change in ("remove", "destination"):
                document = Converter(parser.root, []).convert()
                link = document.paragraphs[0].hyperlinks[0]._hyperlink
                if change == "remove":
                    position = document.paragraphs[0]._p.index(link)
                    for run in list(link):
                        document.paragraphs[0]._p.insert(position, run)
                        position += 1
                    link.getparent().remove(link)
                else:
                    document.part.rels[link.get(qn("r:id"))]._target = "https://wrong.example/"
                document.save(output)
                with self.assertRaisesRegex(ValueError, "Hyperlink"):
                    validate_output(parser.root, output)

    def test_truncated_and_unclosed_link_fields_fail(self):
        parser = parser_with_styles()
        text = control_text(b"%hlk", 3) + b"X\0\r\0"
        for payload in (b"klh%", b"klh%" + struct.pack("<IBH", 0, 0, 999)):
            with self.assertRaisesRegex(ValueError, "Truncated"):
                parser.paragraph(record_tree(paragraph_records(text) + record(71, payload, 1))[0])
        data = link_paragraph().replace(control_text(bytes(4), 4), "x".encode("utf-16le") * 8, 1)
        with self.assertRaisesRegex(ValueError, "missing an end marker"):
            parser.paragraph(record_tree(data)[0])


class HeaderFooterTests(unittest.TestCase):
    def test_binary_header_and_footer_links_use_their_own_parts(self):
        body = (paragraph_records(b"B\0" + control_text(b"head", 16) + control_text(b"foot", 16) + b"\r\0")
                + header_records(contents=link_paragraph(level=2)) + header_records(b"foot", "Footer"))
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "header.hwp", Path(directory) / "header.docx"
            source.write_bytes(hwp_fixture(1, compressed=True, section_bodies=[body]))
            root, warnings = read_hwp(source)
            Converter(root, warnings).convert().save(output)
            document = Document(output)
            self.assertEqual(document.paragraphs[0].text, "B")
            self.assertEqual(document.sections[0].header.paragraphs[0].hyperlinks[0].address,
                             "https://example.org/a;b?x=1&y=2")
            self.assertEqual(document.sections[0].footer.paragraphs[0].text, "Footer")
            report = validate_output(root, output)
            self.assertEqual((report["header_footer_stories"], report["hyperlinks"]), (2, 1))

    def test_odd_even_inheritance_and_explicit_empty_override(self):
        root, columns = sample_root()
        first = paragraph("First")
        for kind, selection, text in (("Header", "odd", "Odd"), ("Header", "even", "Even"),
                                      ("Footer", "both", "Both")):
            story = ET.SubElement(first.find("LineSeg"), kind, {"apply": selection})
            story.append(paragraph(text))
        columns.append(first)
        second = ET.SubElement(root.find("BodyText"), "SectionDef")
        ET.SubElement(second, "ColumnSet").append(paragraph("Second"))
        third = ET.SubElement(root.find("BodyText"), "SectionDef")
        third_body = paragraph("Third")
        ET.SubElement(third_body.find("LineSeg"), "Header", {"apply": "both"})
        ET.SubElement(third, "ColumnSet").append(third_body)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "headers.docx"
            Converter(root, []).convert().save(output)
            document = Document(output)
            self.assertTrue(document.settings.odd_and_even_pages_header_footer)
            self.assertEqual(document.sections[0].header.paragraphs[0].text, "Odd")
            self.assertEqual(document.sections[0].even_page_header.paragraphs[0].text, "Even")
            self.assertTrue(document.sections[1].header.is_linked_to_previous)
            self.assertEqual(document.sections[1].header.paragraphs[0].text, "Odd")
            self.assertEqual(document.sections[1].even_page_footer.paragraphs[0].text, "Both")
            self.assertEqual(document.sections[2].header.paragraphs[0].text, "")
            self.assertEqual(document.sections[2].even_page_header.paragraphs[0].text, "")
            self.assertTrue(validate_output(root, output)["header_footer_content_preserved"])
            inherited = Document(output)
            inherited.sections[1].header.is_linked_to_previous = False
            inherited.sections[1].header.paragraphs[0].text = "Unexpected override"
            inherited.save(output)
            with self.assertRaisesRegex(ValueError, "inheritance"):
                validate_output(root, output)
            document.sections[0].header.paragraphs[0].text = "Changed"
            document.save(output)
            with self.assertRaisesRegex(ValueError, "Text preservation"):
                validate_output(root, output)

    def test_table_in_header(self):
        contents = (paragraph_records(control_text(b"tbl ") + b"\r\0", 2)
                    + table_records("Cell\r".encode("utf-16le"), level=3))
        data = (paragraph_records(control_text(b"head", 16) + b"\r\0")
                + header_records(contents=contents))
        parser = parser_with_styles()
        parser.section.find("ColumnSet").append(parser.paragraph(record_tree(data)[0]))
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "table.docx"
            Converter(parser.root, []).convert().save(output)
            self.assertEqual(Document(output).sections[0].header.tables[0].cell(0, 0).text, "Cell")
            self.assertEqual(validate_output(parser.root, output)["tables"], 1)

    def test_duplicate_or_late_populated_definitions_are_rejected(self):
        root, columns = sample_root()
        host = paragraph("Body")
        header = ET.SubElement(host.find("LineSeg"), "Header")
        header.append(paragraph("Header"))
        columns.append(host)
        duplicate = deepcopy(header)
        host.find("LineSeg").append(duplicate)
        with self.assertRaisesRegex(ValueError, "Repeated"):
            Converter(root, []).convert()
        host.find("LineSeg").remove(duplicate)
        columns.insert(0, paragraph("Earlier page"))
        with self.assertRaisesRegex(ValueError, "first section paragraph"):
            Converter(root, []).convert()


class ImageTransportTests(unittest.TestCase):
    def test_worker_spools_images_and_cli_removes_temporary_files(self):
        picture = paragraph_records(control_text(b"gso ") + b"\r\0", 2) + picture_records(3)
        body = paragraph_records(control_text(b"head", 16) + b"\r\0") + header_records(contents=picture)
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "image.hwp", Path(directory) / "image.docx"
            source.write_bytes(hwp_fixture(1, section_bodies=[body], assets=[("BIN0001.png", PNG)]))
            assets = Path(directory) / "assets"
            assets.mkdir()
            root, warnings = read_hwp(source, asset_directory=assets)
            asset = root.find("BinDataEmbedding")
            self.assertIsNone(asset.text)
            self.assertEqual(asset.get("file"), "BIN0001.png")
            self.assertTrue((assets / asset.get("file")).read_bytes().startswith(PNG))
            Converter(root, warnings, asset_directory=assets).convert().save(output)
            self.assertEqual(validate_output(root, output)["images"], 1)
            cli_output = Path(directory) / "cli.docx"
            with patch("sys.argv", ["convert_hwp.py", str(source), str(cli_output)]), \
                    patch("convert_hwp.read_hwp", wraps=read_hwp) as reader, redirect_stdout(StringIO()):
                main()
            self.assertFalse(Path(reader.call_args.kwargs["asset_directory"]).exists())
            self.assertEqual(json.loads(cli_output.with_suffix(".report.json").read_text())["images"], 1)
            bad_map = Path(directory) / "bad-fonts.json"
            bad_map.write_text("[]")
            failed_output = Path(directory) / "failed.docx"
            with patch("sys.argv", ["convert_hwp.py", str(source), str(failed_output), "--font-map", str(bad_map)]), \
                    patch("convert_hwp.read_hwp", wraps=read_hwp) as reader, \
                    patch("sys.stderr", StringIO()), self.assertRaises(SystemExit):
                main()
            self.assertFalse(Path(reader.call_args.kwargs["asset_directory"]).exists())
            self.assertFalse(failed_output.exists())

    def test_extracted_image_paths_cannot_escape_the_asset_directory(self):
        root, columns = sample_root()
        host = paragraph("Image")
        drawing = ET.SubElement(host.find("LineSeg"), "GShapeObjectControl", {"width": "7200", "height": "7200"})
        picture = ET.SubElement(drawing, "ShapePicture")
        ET.SubElement(picture, "PictureInfo", {"bindata-id": "1"})
        ET.SubElement(root, "BinDataEmbedding", {"storage-id": "BIN0001", "file": "../outside.png"})
        columns.append(host)
        with tempfile.TemporaryDirectory() as directory, self.assertRaisesRegex(ValueError, "image path"):
            Converter(root, [], asset_directory=directory).convert()


if __name__ == "__main__":
    unittest.main()
