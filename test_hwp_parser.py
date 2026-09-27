from io import BytesIO
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zlib

from docx import Document

from convert_hwp import Converter, validate_output
from hwp_parser import HwpParser, decompress_stream, element, record_tree, records


def record(tag, payload=b"", level=0, extended=False):
    size = 4095 if extended else len(payload)
    header = struct.pack("<I", tag | (level << 10) | (size << 20))
    return header + (struct.pack("<I", len(payload)) if extended else b"") + payload


def paragraph_records(text, level=0):
    header = struct.pack("<IIHBBHHHI", len(text) // 2, 0, 0, 0, 0, 1, 0, 1, 0)
    return (record(66, header, level) + record(67, text, level + 1)
            + record(68, struct.pack("<II", 0, 0), level + 1)
            + record(69, struct.pack("<I7iI", 0, 0, 1000, 1000, 800, 200, 0, 7200, 0), level + 1))


def control_text(identifier, code=11):
    return struct.pack("<H4sIIH", code, identifier[::-1], 0, 0, code)


def table_records(text, level=1, colspan=1, nested=b""):
    control = record(71, struct.pack("<4s5I", b" lbt", 1, 0, 0, 7200, 2000), level)
    definition = record(77, struct.pack("<I7H2H", 0, 1, 1, 0, 0, 0, 0, 0, 1, 0), level + 1)
    cell = record(72, struct.pack("<II4H2I5H", 1, 0, 0, 0, colspan, 1, 7200, 2000, 0, 0, 0, 0, 0), level + 1)
    return control + definition + cell + paragraph_records(text, level + 1) + nested


def parser_with_styles(recover_text=False):
    parser = HwpParser(recover_text=recover_text)
    info = ET.SubElement(parser.root, "DocInfo")
    ET.SubElement(info, "CharShape", {"basesize": "1000"})
    ET.SubElement(info, "ParaShape", {"linespacing-type": "ratio"})
    body = ET.SubElement(parser.root, "BodyText")
    parser.section = ET.SubElement(body, "SectionDef")
    ET.SubElement(parser.section, "ColumnSet")
    return parser


class RecordTests(unittest.TestCase):
    def test_regular_extended_and_nested_records(self):
        parsed = list(records(record(66, b"parent") + record(67, b"child", 1, True)))
        self.assertEqual([(item.tag, item.level, item.data) for item in parsed],
                         [(66, 0, b"parent"), (67, 1, b"child")])

    def test_rejects_truncation_and_invalid_nesting(self):
        for data in (b"\0", struct.pack("<I", 4095 << 20), record(66, b"abc")[:-1],
                     record(66, level=1), record(66) + record(67, level=2)):
            with self.subTest(data=data), self.assertRaises(ValueError):
                list(records(data))

    def test_bounded_raw_deflate(self):
        compressor = zlib.compressobj(wbits=-15)
        data = compressor.compress(b"hello" * 100) + compressor.flush()
        self.assertEqual(decompress_stream(data), b"hello" * 100)
        trailer = struct.pack("<II", zlib.crc32(b"hello" * 100), 500)
        self.assertEqual(decompress_stream(data + trailer), b"hello" * 100)
        for invalid in (data[:-1], data + b"trailing", b"bad"):
            with self.subTest(data=invalid), self.assertRaises(ValueError):
                decompress_stream(invalid)
        with self.assertRaisesRegex(ValueError, "size limit"):
            decompress_stream(data, limit=20)

    def test_document_fonts_character_and_paragraph_styles(self):
        parser = HwpParser()
        mappings = record(17, struct.pack("<8I", 0, *([1] * 7)))
        name = "Test Font".encode("utf-16le")
        fonts = record(19, struct.pack("<BH", 0, len(name) // 2) + name, 1) * 7
        shape = struct.pack("<7H7B7b7B7biIbb4I", *([0] * 7), *([100] * 7),
                            *([-3] * 7), *([100] * 7), *([0] * 7), 2000, 2, 0, 0,
                            0x00332211, 0, 0, 0)
        paragraph_shape = struct.pack("<I6iH", 3 << 2, 100, 200, -300, 400, 500, 160, 0)
        parser.doc_info(mappings + fonts + record(21, shape, 1) + record(25, paragraph_shape, 1))
        char = parser.root.find("DocInfo/CharShape")
        self.assertEqual(char.get("bold"), "1")
        self.assertEqual(char.get("text-color"), "#112233")
        self.assertEqual(char.find("LetterSpacing").get("ko"), "-3")
        para = parser.root.find("DocInfo/ParaShape")
        self.assertEqual(para.get("align"), "center")
        self.assertEqual(para.get("indent"), "-300")


class NativeDocumentTests(unittest.TestCase):
    def test_unicode_controls_and_order(self):
        parser = parser_with_styles()
        text = ("Before \U0001f600".encode("utf-16le") + control_text(b"\0\0\0\0", 9)
                + "After\nEnd\r".encode("utf-16le"))
        paragraph = parser.paragraph(record_tree(paragraph_records(text))[0])
        self.assertEqual("".join(item.text for item in paragraph.iter("Text")), "Before \U0001f600AfterEnd")
        self.assertEqual([item.get("name") for item in paragraph.iter("ControlChar")],
                         ["TAB", "LINE_BREAK", "PARAGRAPH_BREAK"])

    def test_nested_tables_are_decoded_and_converted(self):
        parser = parser_with_styles()
        marker = control_text(b"tbl ") + b"\r\0"
        data = paragraph_records(marker) + table_records(marker, nested=table_records("Inner\r".encode("utf-16le"), level=3))
        parser.section.find("ColumnSet").append(parser.paragraph(record_tree(data)[0]))
        document = Converter(parser.root, []).convert()
        self.assertEqual(document.tables[0].cell(0, 0).tables[0].cell(0, 0).text, "Inner")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "nested.docx"
            document.save(output)
            self.assertEqual(validate_output(parser.root, output)["tables"], 2)

    def test_table_span_outside_grid_is_rejected(self):
        parser = parser_with_styles()
        data = paragraph_records(control_text(b"tbl ") + b"\r\0") + table_records(b"\r\0", colspan=2)
        with self.assertRaisesRegex(ValueError, "cell span"):
            parser.paragraph(record_tree(data)[0])

    def test_missing_control_and_unknown_control_are_rejected(self):
        for control in (b"", record(71, b"xxxx", 1)):
            parser = parser_with_styles()
            data = paragraph_records(control_text(b"xxxx") + b"\r\0") + control
            with self.subTest(control=control), self.assertRaises(ValueError):
                parser.paragraph(record_tree(data)[0])

    def test_truncated_control_is_rejected(self):
        parser = parser_with_styles()
        with self.assertRaisesRegex(ValueError, "Truncated paragraph control"):
            parser.paragraph(record_tree(paragraph_records(b"\x0b\0\r\0"))[0])

    def test_recovery_is_explicit_and_reports_invalid_text(self):
        data = paragraph_records(b"\x00\xd8" + "Text\r".encode("utf-16le"))
        with self.assertRaisesRegex(ValueError, "Invalid UTF-16"):
            parser_with_styles().paragraph(record_tree(data)[0])
        parser = parser_with_styles(recover_text=True)
        paragraph = parser.paragraph(record_tree(data)[0])
        self.assertEqual("".join(item.text for item in paragraph.iter("Text")), "\ufffdText")
        self.assertEqual(parser.recovery.replacements[0]["invalid_bytes_hex"], "00d8")

    def test_populated_header_is_preserved(self):
        parser = parser_with_styles()
        data = (paragraph_records(control_text(b"head", 16) + b"\r\0")
                + record(71, b"daeh" + bytes(4), 1) + paragraph_records("Header\r".encode("utf-16le"), 2))
        parser.section.find("ColumnSet").append(parser.paragraph(record_tree(data)[0]))
        document = Converter(parser.root, []).convert()
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "Header")

    def test_picture_reference_and_docx_embedding(self):
        parser = parser_with_styles()
        data = (paragraph_records(control_text(b"gso ") + b"\r\0")
                + record(71, struct.pack("<4s5I", b" osg", 1, 0, 0, 7200, 3600), 1)
                + record(76, b"cip$", 2)
                + record(85, bytes(71) + struct.pack("<H", 1) + bytes(5), 3))
        parser.section.find("ColumnSet").append(parser.paragraph(record_tree(data)[0]))
        asset = element("BinDataEmbedding", storage_id="BIN0001", ext="png", inline="true")
        asset.text = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a1XkAAAAASUVORK5CYII="
        parser.root.append(asset)
        document = Converter(parser.root, []).convert()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "picture.docx"
            document.save(output)
            self.assertEqual(validate_output(parser.root, output)["images"], 1)

    def test_textbox_preserves_controls_and_empty_paragraphs(self):
        parser = parser_with_styles()
        text = ("First".encode("utf-16le") + control_text(b"\0\0\0\0", 9)
                + "Second\nThird\x1fFourth\r".encode("utf-16le"))
        data = (paragraph_records(control_text(b"gso ") + b"\r\0")
                + record(71, struct.pack("<4s5I", b" osg", 1, 0, 0, 7200, 3600), 1)
                + record(76, b"cer$", 2) + record(72, struct.pack("<II", 3, 0), 3)
                + paragraph_records(text, 4) + paragraph_records(b"\r\0", 4)
                + paragraph_records("Fifth\r".encode("utf-16le"), 4))
        parser.section.find("ColumnSet").append(parser.paragraph(record_tree(data)[0]))
        document = Converter(parser.root, []).convert()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "textbox.docx"
            document.save(output)
            self.assertEqual(Document(output).paragraphs[0].text,
                             "First\tSecond\nThird\u2007Fourth\n\nFifth")
            self.assertTrue(validate_output(parser.root, output)["text_preserved"])

    def test_picture_before_table_is_not_deleted(self):
        parser = parser_with_styles()
        data = (paragraph_records(control_text(b"gso ") + control_text(b"tbl ") + b"\r\0")
                + record(71, struct.pack("<4s5I", b" osg", 1, 0, 0, 7200, 3600), 1)
                + record(76, b"cip$", 2)
                + record(85, bytes(71) + struct.pack("<H", 1) + bytes(5), 3)
                + table_records("Cell\r".encode("utf-16le")))
        parser.section.find("ColumnSet").append(parser.paragraph(record_tree(data)[0]))
        asset = element("BinDataEmbedding", storage_id="BIN0001", ext="png", inline="true")
        asset.text = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a1XkAAAAASUVORK5CYII="
        parser.root.append(asset)
        document = Converter(parser.root, []).convert()
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "picture-table.docx"
            document.save(output)
            report = validate_output(parser.root, output)
            self.assertEqual((report["images"], report["tables"]), (1, 1))
            self.assertEqual(len(document.tables[0]._tbl.getprevious().xpath(".//w:drawing")), 1)

    def test_native_reader_rejects_protection_before_body_read(self):
        for flag in (2, 4, 16, 256, 1024):
            with self.subTest(flag=flag), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "protected.hwp"
                source.touch()
                header = b"HWP Document File".ljust(32, b"\0") + bytes([0, 0, 0, 5]) + struct.pack("<I", flag)
                with patch("hwp_parser.olefile.OleFileIO") as container:
                    document = container.return_value.__enter__.return_value
                    document.get_size.return_value = len(header)
                    document.openstream.return_value = BytesIO(header)
                    with self.assertRaisesRegex(ValueError, "protected"):
                        HwpParser().parse(source)
                    document.openstream.assert_called_once_with("FileHeader")

    def test_stream_compression_modes(self):
        payload = b"Native HWP stream"
        compressor = zlib.compressobj(wbits=-15)
        compressed = compressor.compress(payload) + compressor.flush()
        for data, is_compressed in ((payload, False), (compressed, True)):
            with self.subTest(compressed=is_compressed), patch("hwp_parser.olefile.OleFileIO") as container:
                container.get_size.return_value = len(data)
                container.openstream.return_value = BytesIO(data)
                self.assertEqual(HwpParser().stream(container, "DocInfo", is_compressed), payload)

    def test_native_sample_matches_saved_text_baselines(self):
        samples = (
            ("2025 멘토링 프로그램_하반기.hwp", 740, "f7a3d0ee230228156b3df4e3128f21aedc7c801f16b17e151e239c4e2b037951"),
            ("5. 협력업체 업무협약서.hwp", 1487, "4867a7dab6748e38786d0eb894570d9031555a5c1e79e450a36e8380872ed6fe"),
        )
        for name, count, digest in samples:
            source = Path("output") / name
            if not source.exists():
                self.skipTest("Local private regression samples are unavailable.")
            with self.subTest(sample=name), tempfile.TemporaryDirectory() as directory:
                root, warnings = HwpParser().parse(source)
                output = Path(directory) / "sample.docx"
                Converter(root, warnings).convert().save(output)
                report = validate_output(root, output)
                self.assertEqual(report["text_characters"], count)
                self.assertEqual(report["text_sha256"], digest)
                self.assertEqual(report["tables"], 2)


if __name__ == "__main__":
    unittest.main()
